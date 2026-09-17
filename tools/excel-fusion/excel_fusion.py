#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Klasör veya kaynak sayfa bazlı, sezgisel başlık eşleştirmeli Excel raporlama.

Python 3.11+
Kurulum: python -m pip install -r requirements.txt
Çalıştırma: python excel_rapor.py --input "C:\\Veriler" --output "C:\\Raporlar\\Ana.xlsx"

Kaynak dosyalar DEĞİŞTİRİLMEZ. Çıktı, verilerin rapor üretim anındaki görüntüsüdür.
Konfigürasyon gerçek bir GUI değildir; aşağıdaki açıklamalı düzenleme alanıdır.
"""
from __future__ import annotations

import argparse
import fnmatch
import itertools
import json
import logging
import math
import os
import re
import sys
import tempfile
import unicodedata
import zipfile
from collections import Counter, deque
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence

# ============================================================================
# 1. KONFİGÜRASYON PANELİ — Normalde yalnızca bu bölümü değiştirmeniz yeterlidir.
# ============================================================================
SCRIPT_DIR = Path(__file__).resolve().parent

CONFIG: dict[str, Any] = {
    # Göreli yollar terminale değil, BU .py DOSYASININ klasörüne göre çözülür.
    # Örnek: Path(r"C:\Users\yasar\Desktop\ExcelKaynaklari")
    "INPUT_DIR": SCRIPT_DIR / "kaynaklar",
    "OUTPUT_FILE": SCRIPT_DIR / "ciktilar" / "Ana_Rapor.xlsx",
    "OVERWRITE": False,                 # True: mevcut çıktı raporunu yeniler.

    # Ana rapor sayfalarının hangi kaynağa göre gruplanacağı:
    # subfolder: mevcut davranış; kök altındaki ilk düzey klasör -> rapor sayfası.
    # sheet_name: tüm alt klasörler ve kökteki Excel'ler taranır; aynı adlı kaynak
    #             sayfalar farklı dosyalarda olsalar da tek rapor sayfasında birleştirilir.
    "REPORT_GROUP_MODE": "subfolder",

    # Kaynaklarda beklenen başlık satırı. Satır numaraları Excel gibi 1'den başlar.
    "EXPECTED_HEADER_ROW": 12,
    "HEADER_SEARCH_START": 1,           # Kaymış başlıklar için tarama başlangıcı.
    "HEADER_SEARCH_END": 40,            # Örnek: başlık 55'e kayabiliyorsa 70 yapın.
    "HEADER_HEIGHTS": (1, 2),           # Bir veya iki AYRI Excel satırındaki başlık.

    # AYNI HÜCREDE: "Reviewer Name\nİnceleyenin adını yazınız" -> "Reviewer Name".
    # smart: ilk dolu satırı esas alır; "Reviewer\nName" gibi bölünmüş adları da arar.
    # first_line: kesin olarak SADECE ilk dolu satırı kullanır; devamını hiç aramaz.
    # full: eski davranış; hücrenin tüm metni birlikte karşılaştırılır.
    "HEADER_TEXT_MODE": "smart",
    # smart modunda başlığın en çok kaç başlangıç satırı birleştirilebilir?
    # 3: "Name\nof\nReviewer\nAçıklama..." için ilk üç satırı da değerlendirir.
    # Sadece baştan başlayan parçalar denenir; açıklamanın içinden ad aranmaz.
    "HEADER_MAX_PREFIX_LINES": 3,
    "MIN_HEADER_MATCHES": 3,            # En az kaç TANIMLI sütun bulunmalı?
    "HEADER_ANCHORS": ("No.", "Entity", "Panel", "ATA", "Reviewer Name"),
    "MIN_HEADER_ANCHORS": 2,            # Yukarıdaki anahtarların en az ikisi.

    # 0–100 metin benzerliği puanı; olasılık/güven yüzdesi değildir.
    "FUZZY_THRESHOLD": 84.0,            # Düşürmek daha esnek ama daha risklidir.
    "AMBIGUITY_MARGIN": 7.0,            # En iyi iki eşleşme arasında gereken fark.
    "SHORT_ALIAS_MAX_LENGTH": 3,        # ATA/No gibi kısa adlarda yalnızca tam eşleşme.

    # Eşleşmeyen başlıklar "Ek | ..." olarak KORUNUR; başka sütuna zorlanmaz.
    # Aynı ek başlık, harf/noktalama normalizasyonuyla aynı sütunda toplanır.
    # Bilinmeyen sütunlarda fuzzy birleştirme YAPILMAZ; bunun için COLUMN_SCHEMA'ya ekleyin.
    "EXTRA_PREFIX": "Ek | ",

    # Yalnızca subfolder modunda: kökteki Excel'leri ayrı bir grupta toplar.
    # sheet_name modunda kökteki Excel'ler bu ayardan bağımsız olarak taranır.
    "INCLUDE_ROOT_FILES": False,
    "ROOT_GROUP_NAME": "Kök Dosyalar",
    "EXTENSIONS": (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".xlsb"),
    "EXCLUDE_DIRS": (".git", ".venv", "venv", "__pycache__"),
    "EXCLUDE_FILES": ("~$*", ".~*", "*.tmp.xlsx"),
    "EXCLUDE_SHEETS": (),              # Örnek: ("*Summary*", "*Dashboard*", "Kapak")
    "INCLUDE_HIDDEN_SHEETS": True,      # Varsayılan: gizli veri sayfaları da taranır.

    # Veriyi sessizce silmemek için elenen DOLU satırlar Karantina'ya kaydedilir.
    # Bağımsız ve öncelikli kontrol: listedeki HER sütun dolu olmalıdır.
    # Kaynak başlığı veya standart çıktı adı; COLUMN_SCHEMA'da bulunması gerekmez.
    # Eksik başlık da boş kabul edilir. Dolgudan önce uygulanır; () kontrolü kapatır.
    "ROW_REQUIRED_COLUMNS": (),        # Örnek: ("Entity", "Notes")
    "MIN_ROW_VALUES": 2,               # Bir veri satırında en az iki dolu hücre.
    "ROW_REQUIRE_ANY": ("No.", "Entity", "Panel", "ATA", "Reviewer Name"),
    # None: sütun değeri yerine kaynak Excel satır numarası sanal kayıt ID'sidir.
    "ROW_ID_COLUMN": "No.",
    "ROW_ID_REGEX": None,              # Örnek: r"^\d+(?:\.\d+)?$"; None: ID zorlanmaz.
    "FOOTER_LABELS": ("total", "grand total", "toplam", "genel toplam",
                      "signature", "imza", "prepared by", "hazirlayan",
                      "approved by", "onaylayan"),
    "TRIM_TEXT": True,                 # Metnin başı/sonu temizlenir; içi değiştirilmez.
    "FILL_DOWN_COLUMNS": (),           # Örnek: ("Entity", "Panel"); boş hücreyi üstten alır.
                                        # Varsayılan boş: birleştirilmiş hücrelerde tahmin yok.
    "DROP_DUPLICATES": False,          # Varsayılan: tekrar gibi görünen kayıtlar bile korunur.
    "DEDUP_KEYS": (),                  # Boşsa tüm veri alanları; kaynak bilgisi hariç.
                                        # Örnek: ("Entity", "No.", "Reviewer Name")

    # Eksik formül önbelleğinde formül metni literal olarak saklanır ve uyarı verilir.
    # Formüller yeni konumda yeniden çalıştırılmaz; yanlış referans üretilmez.
    "PRESERVE_SIMPLE_ZERO_FORMATS": True,  # Sayısal 7 + '000' formatı -> metin '007'.

    "REPORT_TITLE": "BİRLEŞTİRİLMİŞ EXCEL RAPORU",
    "SUMMARY_SHEET": "Özet & Dashboard",
    "DASHBOARD_TOP_N": 10,
    # Açık/kapalı ve iş sırası özeti; yalnızca rapora alınan kayıtlar sayılır.
    "WORKFLOW_SUMMARY": {
        "ENABLED": True,               # False: bu özet bölümünü gösterme.
        # COLUMN_SCHEMA'daki standart adlar; kaynak alternatifleri aliases ile tanımlanır.
        "STATUS_COLUMN": "Open / Closed",
        "COMPANY_COLUMN": "My Company",
        # Baş/son boşluk ve harf büyüklüğü yok sayılır; devamındaki metin serbesttir.
        "OPEN_PREFIX": "open",
        "CLOSED_PREFIX": "closed",
        # Yalnızca açık kayıtta: My Company boşsa bizde, doluysa diğer şirkette.
        # Eksik sütun/okunamayan formül ile bilinmeyen durum ayrı sayılır.
    },
    # Başlık adı COLUMN_SCHEMA'daki standart ad olmalıdır. Klasör özel boyuttur;
    # sheet_name modunda bu boyut kaynak sayfa gruplarını gösterir.
    "CHARTS": (("Klasör", "bar"), ("ATA", "column"),
               ("Panel", "doughnut"), ("Reviewer Name", "bar")),
    "MAX_ROWS_PER_SHEET": 1_048_000,     # Büyük gruplar Grup, Grup (2)... diye bölünür.
    "DATE_FORMAT": "dd.mm.yyyy",
    "DATETIME_FORMAT": "dd.mm.yyyy hh:mm",
    "FAIL_ON_FILE_ERROR": False,        # True: tek bir dosya hatasında çıktı üretilmez.
}


@dataclass(frozen=True)
class ColumnSpec:
    """name: çıktı adı; aliases: alternatifler; required: başlıkta zorunlu alan.
    threshold: bu alana özel 0–100 eşik (None: genel eşik); width: Excel genişliği.
    """
    name: str
    aliases: tuple[str, ...] = ()
    required: bool = False
    threshold: float | None = None
    width: float = 24


# Çıktının standart sütun sırası BU LİSTENİN sırasıdır; kaynak sırası önemli değildir.
# YENİ SÜTUNLARI aşağıdaki örnekleri kopyalayarak ekleyebilirsiniz.
COLUMN_SCHEMA = [
    ColumnSpec("No.", ("No", "Number", "Item No", "Item Number", "Sıra No",
                       "Sıra Numarası", "Numara"), width=11),
    ColumnSpec("Entity", ("Entitiy", "Entiy", "Varlık", "Kurum"), width=27),
    ColumnSpec("Panel", ("Panel Name", "Panel Adı", "Pannel"), width=22),
    ColumnSpec("ATA", ("ATA No", "ATA Number", "ATA Chapter", "ATA Kodu"), width=14),
    ColumnSpec("Reviewer Name", ("Reviever Name", "Rewiever Name", "Reviwer Name",
                                "Name of Reviewer", "Reviewer", "İnceleyen",
                                "İnceleyen Adı", "Değerlendiren"), width=28),
    ColumnSpec("Open / Closed", ("Open Closed", "Open/Closed"), width=30),
    ColumnSpec("My Company", width=30),
    # ColumnSpec("Comment", ("Comments", "Review Comment", "Açıklama", "Yorum"), width=55),
    # ColumnSpec("Status", ("Staus", "Review Status", "Durum"), width=20),
    # ColumnSpec("Due Date", ("Deadline", "Target Date", "Termin", "Bitiş Tarihi"), width=19),
    # ColumnSpec("Priority", ("Öncelik", "Prioriy"), width=18),
    # ColumnSpec("Document No", ("Document Number", "Doküman No"), required=True, width=28),
]
# required=True olan sütun bulunamazsa sayfa veri olarak alınmaz ve denetime yazılır.
# "No" alternatifini hem No. hem Document No alanına eklemeyin: çakışma hata verir.

PALETTE = {
    "navy": "#14243B", "teal": "#0EAB9D", "cyan": "#28BCE0",
    "violet": "#8260DB", "amber": "#F0AD32", "rose": "#E86783",
    "paper": "#F2F6FB", "ink": "#263449", "muted": "#6B7B90",
    "line": "#DFE7F1", "white": "#FFFFFF",
}
# ======================= KONFİGÜRASYON SONU ================================

LOG = logging.getLogger("excel_rapor")
REPORT_MARKER = "FolderExcelReport-v1"
REPORT_GROUP_MODES = ("subfolder", "sheet_name")
META_COLUMNS = ("Kaynak Dosya", "Kaynak Sayfa", "Kaynak Satır")
MISSING_FORMULA_PREFIX = "[FORMÜL SONUCU OKUNAMADI]"
AUX_SHEETS = ("Denetim", "Başlık Eşleştirmeleri", "Karantina", "_Rapor Verisi")


@dataclass
class SourceRow:
    number: int
    values: list[Any]
    formats: list[str] = field(default_factory=list)
    missing_formulas: list[tuple[int, str]] = field(default_factory=list)


@dataclass
class Match:
    index: int
    raw: str
    target: str | None
    score: float
    state: str
    alternative: str = ""
    margin: float = 100.0
    label: str = ""                    # Açıklamadan ayıklanıp karşılaştırılan metin.
    strategy: str = ""                 # Denetimde seçimin nasıl yapıldığını gösterir.
    scored_label: str = ""             # Skorun hesaplandığı aday (eşleşmeyenlerde farklı olabilir).


@dataclass
class Header:
    start: int
    height: int
    labels: list[str]
    matches: list[Match]
    quality: float


@dataclass
class Record:
    values: dict[str, Any]
    file: str
    sheet: str
    row: int


@dataclass
class Group:
    name: str
    files: list[Path] = field(default_factory=list)
    records: list[Record] = field(default_factory=list)
    extra_names: dict[tuple[str, int], str] = field(default_factory=dict)
    read_files: set[str] = field(default_factory=set)
    table_count: int = 0
    sheet_names: list[str] = field(default_factory=list)
    seen: set[str] = field(default_factory=set)


@dataclass
class Report:
    root: Path
    output: Path
    groups: list[Group] = field(default_factory=list)
    source_files: list[Path] = field(default_factory=list)
    events: list[dict[str, Any]] = field(default_factory=list)
    mappings: list[dict[str, Any]] = field(default_factory=list)
    quarantine: list[dict[str, Any]] = field(default_factory=list)
    generated_at: datetime = field(default_factory=datetime.now)

    def event(self, level: str, code: str, group: str = "", file: str = "",
              sheet: str = "", row: int | None = None, detail: str = "") -> None:
        self.events.append(dict(Seviye=level, Kod=code, Klasör=group,
                                Dosya=file, Sayfa=sheet, Satır=row, Açıklama=detail))
        getattr(LOG, {"ERROR": "error", "WARNING": "warning"}.get(level, "info"))(
            "%s | %s | %s | %s", code, file, sheet, detail)


def nonempty(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def usable_value(value: Any) -> bool:
    return nonempty(value) and not (isinstance(value, str) and value.startswith(MISSING_FORMULA_PREFIX))


@lru_cache(maxsize=32768)
def normalize(text: str) -> str:
    """Türkçe harf, aksan, satır sonu, boşluk ve noktalama farklarını azaltır."""
    text = text.replace("ı", "i").replace("İ", "I").replace("\u00ad", "")
    text = unicodedata.normalize("NFKD", text).casefold()
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^\w]+", " ", text, flags=re.UNICODE).replace("_", " ").split())


@lru_cache(maxsize=65536)
def similarity(a: str, b: str, short_limit: int) -> float:
    """Kısa anahtarlar tam eşleşir; uzunlarda karakter ve kelime sırası karşılaştırılır."""
    x, y = a.replace(" ", ""), b.replace(" ", "")
    if not x or not y:
        return 0.0
    if a == b or x == y:
        return 100.0
    if min(len(x), len(y)) <= short_limit:
        return 0.0
    direct = SequenceMatcher(None, x, y, autojunk=False).ratio()
    tokens = SequenceMatcher(None, " ".join(sorted(a.split())),
                             " ".join(sorted(b.split())), autojunk=False).ratio()
    return 100 * max(direct, tokens)


def header_text_candidates(raw: str, cfg: dict[str, Any]) -> list[str]:
    """Sadece başlık için aday çıkarır; veri hücrelerine ASLA uygulanmaz.

    Satır sonları normalizasyondan ÖNCE ayrılır. Yalnızca ilk dolu satırdan
    başlayan önekler denenir; açıklama satırları tek başına başlık sayılmaz.
    """
    mode = cfg.get("HEADER_TEXT_MODE", "smart")
    if mode == "full":
        return [raw.strip()]
    lines = [line.strip() for line in raw.splitlines() if line.strip()]
    if not lines:
        return [""]
    limit = 1 if mode == "first_line" else cfg.get("HEADER_MAX_PREFIX_LINES", 3)
    return [" ".join(lines[:n]) for n in range(1, min(len(lines), limit) + 1)]


class HeaderMatcher:
    def __init__(self, schema: Sequence[ColumnSpec], cfg: dict[str, Any]):
        self.schema, self.cfg = list(schema), cfg
        self.aliases = {s.name: tuple(normalize(a) for a in (s.name, *s.aliases)) for s in schema}

    def _rank(self, text: str) -> list[tuple[float, ColumnSpec]]:
        text = normalize(text)
        return sorted(((max(similarity(text, a, self.cfg["SHORT_ALIAS_MAX_LENGTH"])
                            for a in self.aliases[s.name]), s) for s in self.schema),
                      key=lambda x: (-x[0], x[1].name))

    def _select_label(self, raw: str) -> tuple[str, list[tuple[float, ColumnSpec]], str]:
        labels = header_text_candidates(raw, self.cfg)
        ranks = [self._rank(label) for label in labels]
        selected, ranked = 0, ranks[0]
        # Bilinen en uzun TAM önek, bölünmüş adı tamamlar. Örn. Reviewer + Name.
        exact = [i for i, ranking in enumerate(ranks) if ranking[0][0] == 100]
        if exact:
            selected = exact[-1]
            ranked = ranks[selected]
        else:
            score, spec = ranks[0][0]
            threshold = spec.threshold if spec.threshold is not None else self.cfg["FUZZY_THRESHOLD"]
            # İlk satır eşiği geçiyorsa açıklama onu bozmasın. Belirsizliği de koru.
            if score < threshold and len(labels) > 1:
                # İlk satır tek başına yetmiyorsa diğer başlangıç öneklerini dene.
                # Farklı öneklerin farklı alanlara yakınlığı belirsizlik hesabına girer.
                best_by_name: dict[str, tuple[float, ColumnSpec, int]] = {}
                for i, ranking in enumerate(ranks):
                    for value, candidate in ranking:
                        old = best_by_name.get(candidate.name)
                        if old is None or value > old[0]:
                            best_by_name[candidate.name] = (value, candidate, i)
                ordered = sorted(best_by_name.values(), key=lambda x: (-x[0], x[1].name))
                ranked = [(value, candidate) for value, candidate, _ in ordered]
                selected = ordered[0][2]
        mode = self.cfg.get("HEADER_TEXT_MODE", "smart")
        strategy = "TÜM METİN" if mode == "full" else f"İLK {selected + 1} DOLU SATIR"
        return labels[selected], ranked, strategy

    def _match_label(self, index: int, raw: str) -> Match:
        label, ranked, strategy = self._select_label(raw)
        scored_label = label
        best, spec = ranked[0]
        second, runner = ranked[1] if len(ranked) > 1 else (0.0, None)
        threshold = spec.threshold if spec.threshold is not None else self.cfg["FUZZY_THRESHOLD"]
        target, state = None, "EŞLEŞMEDİ"
        if best >= threshold:
            if best - second < self.cfg["AMBIGUITY_MARGIN"]:
                state = "BELİRSİZ"
            else:
                target, state = spec.name, "TAM" if best == 100 else "SEZGİSEL"
        # Bilinmeyen alanın uzun açıklaması ayrı ek sütunlar oluşturmasın.
        # Tam ham metin Match.raw ve denetim sayfasında korunur.
        if target is None:
            label = header_text_candidates(raw, self.cfg)[0]
            strategy = "TÜM METİN" if self.cfg.get("HEADER_TEXT_MODE") == "full" else "İLK 1 DOLU SATIR"
        return Match(index, raw, target, round(best, 2), state,
                     runner.name if runner else "", round(best - second, 2), label, strategy, scored_label)

    def match(self, labels: Sequence[str]) -> list[Match]:
        results = [self._match_label(index, str(raw)) for index, raw in enumerate(labels) if nonempty(raw)]
        # Bir hedefe iki farklı kaynak sütunu yazılmaz. Kaybeden sütun ayrı korunur.
        claimed: set[str] = set()
        for item in sorted(results, key=lambda m: (-m.score, m.index)):
            if item.target in claimed:
                item.target, item.state = None, "ÇAKIŞMA"
            elif item.target:
                claimed.add(item.target)
        return sorted(results, key=lambda m: m.index)

    def candidate(self, rows: Sequence[SourceRow], offset: int, height: int) -> Header | None:
        chunk = rows[offset:offset + height]
        if len(chunk) != height:
            return None
        if any(chunk[i].number + 1 != chunk[i + 1].number for i in range(height - 1)):
            return None
        width = max((len(r.values) for r in chunk), default=0)
        # Hücre içi satır sonu ile iki AYRI Excel satırını karıştırmayın.
        # Her hücre önce temizlenir; fiziksel satırlar SONRA birleştirilir.
        row_matches = [self.match([str(v) if nonempty(v) else "" for v in row.values])
                       for row in chunk]
        row_titles = [{m.index: m.label for m in matches} for matches in row_matches]
        raw_labels, labels = [], []
        for col in range(width):
            parts = [str(r.values[col]).strip() for r in chunk
                     if col < len(r.values) and nonempty(r.values[col])]
            raw_labels.append(" ".join(dict.fromkeys(parts)))
            clean_parts = [titles[col] for titles in row_titles if col in titles]
            labels.append(" ".join(dict.fromkeys(clean_parts)))
        matches = row_matches[0] if height == 1 else self.match(labels)
        for item in matches:
            item.raw = raw_labels[item.index]
            if height > 1:
                item.strategy = f"{height} EXCEL SATIRI; HÜCRE İÇİ AYIKLAMA"
        good = [m for m in matches if m.target]
        targets = {m.target for m in good}
        if len(good) < self.cfg["MIN_HEADER_MATCHES"]:
            return None
        anchors = set(self.cfg["HEADER_ANCHORS"])
        if len(targets & anchors) < self.cfg["MIN_HEADER_ANCHORS"]:
            return None
        if any(s.required and s.name not in targets for s in self.schema):
            return None
        if height > 1:
            # İlk veri satırını yutmayın. Yeni alan bulunmalı VEYA alt satır yalnızca
            # "Reviewer" + "Name" gibi tam eşleşen başlık devamlarından oluşmalı.
            single_best = max(sum(m.target is not None for m in self.match(
                [str(v) if nonempty(v) else "" for v in r.values])) for r in chunk)
            split_count, lower_count = 0, 0
            by_index = {m.index: m for m in good}
            for col, lower in enumerate(chunk[1].values):
                if not nonempty(lower):
                    continue
                lower_count += 1
                upper = chunk[0].values[col] if col < len(chunk[0].values) else None
                item = by_index.get(col)
                if isinstance(upper, str) and isinstance(lower, str) and item and item.score == 100:
                    full = normalize(labels[col]).replace(" ", "")
                    upper_title = normalize(row_titles[0].get(col, "")).replace(" ", "")
                    if len(full) > len(upper_title):
                        split_count += 1
            if len(good) <= single_best and not (split_count > 0 and split_count == lower_count):
                return None
        distance = abs(chunk[0].number - self.cfg["EXPECTED_HEADER_ROW"])
        quality = 100 * len(good) + sum(m.score for m in good) / len(good) - 0.2 * distance + 2 * (height - 1)
        return Header(chunk[0].number, height, raw_labels, matches, quality)

    def detect(self, rows: Sequence[SourceRow]) -> Header | None:
        candidates = []
        for i, row in enumerate(rows):
            if not self.cfg["HEADER_SEARCH_START"] <= row.number <= self.cfg["HEADER_SEARCH_END"]:
                continue
            for height in self.cfg["HEADER_HEIGHTS"]:
                result = self.candidate(rows, i, height)
                if result:
                    candidates.append(result)
        # 12. satıra yakın ilk tabloyu tercih edin; aşağıdaki daha temiz bir tekrar
        # başlığı seçip aradaki gerçek kayıtları atlamayın.
        return min(candidates, key=lambda h: (abs(h.start - self.cfg["EXPECTED_HEADER_ROW"]),
                                              -h.quality, h.start), default=None)


def validate(cfg: dict[str, Any], schema: Sequence[ColumnSpec]) -> None:
    names = [s.name for s in schema]
    group_mode = cfg.get("REPORT_GROUP_MODE", "subfolder")
    if group_mode not in REPORT_GROUP_MODES:
        raise ValueError("REPORT_GROUP_MODE: subfolder veya sheet_name olmalı.")
    if not names or len(set(n.casefold() for n in names)) != len(names):
        raise ValueError("COLUMN_SCHEMA boş veya standart sütun adları tekrarlı.")
    reserved = {normalize(n) for n in (*META_COLUMNS, "Alan Doluluğu")}
    if any(normalize(n) in reserved for n in names):
        raise ValueError("Standart sütun adı kaynak/yardımcı sütun adlarıyla çakışıyor.")
    aliases: dict[str, str] = {}
    for spec in schema:
        if not spec.name.strip() or len(spec.name) > 180 or spec.width <= 0:
            raise ValueError(f"Geçersiz sütun tanımı: {spec}")
        if spec.threshold is not None and not 0 < spec.threshold <= 100:
            raise ValueError(f"Geçersiz sütun eşiği: {spec.name}")
        for a in (spec.name, *spec.aliases):
            key = normalize(a).replace(" ", "")
            if not key:
                raise ValueError("Boş alias kullanılamaz.")
            if key in aliases and aliases[key] != spec.name:
                raise ValueError(f"Alias çakışması: {a!r}: {aliases[key]} / {spec.name}")
            aliases[key] = spec.name
    for key in ("HEADER_ANCHORS", "ROW_REQUIRE_ANY", "FILL_DOWN_COLUMNS", "DEDUP_KEYS"):
        missing = set(cfg[key]) - set(names)
        if missing:
            raise ValueError(f"{key} içinde tanımsız sütun var: {missing}")
    required_columns = cfg.get("ROW_REQUIRED_COLUMNS", ())
    if (not isinstance(required_columns, (list, tuple))
            or any(not isinstance(name, str) or not normalize(name) for name in required_columns)):
        raise ValueError("ROW_REQUIRED_COLUMNS dolu sütun adlarından oluşan liste veya tuple olmalı.")
    if not 1 <= cfg["MIN_HEADER_MATCHES"] <= len(names):
        raise ValueError("MIN_HEADER_MATCHES, 1 ile tanımlı sütun sayısı arasında olmalı.")
    if not 0 <= cfg["MIN_HEADER_ANCHORS"] <= len(cfg["HEADER_ANCHORS"]):
        raise ValueError("MIN_HEADER_ANCHORS geçersiz.")
    if not 1 <= cfg["HEADER_SEARCH_START"] <= cfg["HEADER_SEARCH_END"]:
        raise ValueError("Başlık arama aralığı geçersiz.")
    if not cfg["HEADER_HEIGHTS"] or not set(cfg["HEADER_HEIGHTS"]) <= {1, 2}:
        raise ValueError("HEADER_HEIGHTS yalnızca 1 ve/veya 2 içerebilir.")
    if cfg.get("HEADER_TEXT_MODE", "smart") not in ("smart", "first_line", "full"):
        raise ValueError("HEADER_TEXT_MODE: smart, first_line veya full olmalı.")
    limit = cfg.get("HEADER_MAX_PREFIX_LINES", 3)
    if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 10:
        raise ValueError("HEADER_MAX_PREFIX_LINES 1–10 arasında bir tam sayı olmalı.")
    if not 0 < cfg["FUZZY_THRESHOLD"] <= 100 or not 0 <= cfg["AMBIGUITY_MARGIN"] < 100:
        raise ValueError("Fuzzy eşikleri geçersiz.")
    if not 1 <= cfg["MAX_ROWS_PER_SHEET"] <= 1_048_568:
        raise ValueError("MAX_ROWS_PER_SHEET 1–1048568 arasında olmalı.")
    if not 1 <= cfg["DASHBOARD_TOP_N"] <= 30:
        raise ValueError("DASHBOARD_TOP_N 1–30 arasında olmalı.")
    if len(cfg["CHARTS"]) > 4:
        raise ValueError("Dashboard yerleşimi en fazla 4 grafik içerir.")
    for dimension, kind in cfg["CHARTS"]:
        if dimension != "Klasör" and dimension not in names:
            raise ValueError(f"Grafik sütunu COLUMN_SCHEMA'da yok: {dimension}")
        if kind not in ("bar", "column", "doughnut"):
            raise ValueError(f"Desteklenmeyen grafik: {kind}")
    if cfg["ROW_ID_REGEX"]:
        re.compile(cfg["ROW_ID_REGEX"])
    if cfg["ROW_ID_COLUMN"] is not None and cfg["ROW_ID_COLUMN"] not in names:
        raise ValueError("ROW_ID_COLUMN, None veya tanımlı bir standart sütun olmalı.")
    validate_workflow_config(cfg, names)


def validate_workflow_config(cfg: dict[str, Any], names: Sequence[str]) -> None:
    """Etkin iş sırası özetinde hatalı sütun ve çakışan önekleri reddeder."""
    settings = cfg.get("WORKFLOW_SUMMARY", {"ENABLED": False})
    if not isinstance(settings, dict) or not isinstance(settings.get("ENABLED"), bool):
        raise ValueError("WORKFLOW_SUMMARY.ENABLED True veya False olmalı.")
    if not settings["ENABLED"]:
        return
    for key in ("STATUS_COLUMN", "COMPANY_COLUMN"):
        if settings.get(key) not in names:
            raise ValueError(f"WORKFLOW_SUMMARY.{key} tanımlı bir standart sütun olmalı.")
    if settings["STATUS_COLUMN"] == settings["COMPANY_COLUMN"]:
        raise ValueError("Durum ve şirket sütunları farklı olmalı.")
    prefixes = []
    for key in ("OPEN_PREFIX", "CLOSED_PREFIX"):
        value = settings.get(key)
        if not isinstance(value, str) or not value.strip():
            raise ValueError(f"WORKFLOW_SUMMARY.{key} boş olmayan bir metin olmalı.")
        prefixes.append(value.strip().casefold())
    if prefixes[0].startswith(prefixes[1]) or prefixes[1].startswith(prefixes[0]):
        raise ValueError("Açık ve kapalı önekleri birbiriyle çakışmamalı.")


def excel_col(index: int) -> str:
    """0 -> A; 25 -> Z; 26 -> AA."""
    result, index = "", index + 1
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def unique_sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\x00-\x1f\[\]:*?/\\]", "_", name).strip().strip("'") or "Sayfa"
    if base.casefold() == "history":
        base = "History_"
    candidate, number = base[:31].rstrip("'") or "Sayfa", 2
    while candidate.casefold() in used:
        suffix = f" ({number})"
        candidate = base[:31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def source_text(value: Any, number_format: str, cfg: dict[str, Any]) -> Any:
    if isinstance(value, str):
        return value.strip() if cfg["TRIM_TEXT"] else value
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if (cfg["PRESERVE_SIMPLE_ZERO_FORMATS"] and isinstance(value, (int, float))
            and not isinstance(value, bool) and re.fullmatch(r"0{2,20}", number_format or "")
            and float(value).is_integer()):
        return f"{int(value):0{len(number_format)}d}"
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15:
        return str(value)
    return value


def extra_name(group: Group, raw: str, occurrence: int, cfg: dict[str, Any],
               schema: Sequence[ColumnSpec]) -> str:
    key = (normalize(raw), occurrence)
    if key in group.extra_names:
        return group.extra_names[key]
    base = cfg["EXTRA_PREFIX"] + (raw.strip() or "Başlıksız")
    base = base[:180]
    used = {s.name.casefold() for s in schema} | {n.casefold() for n in group.extra_names.values()}
    candidate, i = base, 2
    while candidate.casefold() in used:
        candidate = f"{base} [{i}]"
        i += 1
    group.extra_names[key] = candidate
    return candidate


def register_header(header: Header, group: Group, report: Report, file: str, sheet: str,
                    cfg: dict[str, Any], schema: Sequence[ColumnSpec]) -> dict[int, str]:
    mapping, occurrences = {}, Counter()
    for item in header.matches:
        clean_title = item.label or item.raw
        raw_key = normalize(clean_title)
        occurrences[raw_key] += 1
        target = item.target or extra_name(group, clean_title, occurrences[raw_key], cfg, schema)
        mapping[item.index] = target
        report.mappings.append({"Klasör": group.name, "Dosya": file, "Sayfa": sheet,
                                "Başlık Satırı": header.start, "Yükseklik": header.height,
                                "Kaynak Sütun": excel_col(item.index), "Orijinal Başlık": item.raw,
                                "Ayıklanan Başlık": clean_title, "Başlık İşleme": item.strategy,
                                "Skorlanan Başlık": item.scored_label or item.label or item.raw,
                                "Hedef Sütun": target, "Eşleşme": item.state, "Skor": item.score,
                                "İkinci Aday": item.alternative, "Skor Farkı": item.margin})
        if item.state in ("BELİRSİZ", "ÇAKIŞMA"):
            report.event("WARNING", "HEADER_REVIEW", group.name, file, sheet, header.start,
                         f"{item.raw!r} -> {target!r}; {item.state}; puan={item.score}")
    missing = [s.name for s in schema if s.name not in mapping.values()]
    if missing:
        report.event("WARNING", "MISSING_COLUMNS", group.name, file, sheet, header.start,
                     "Bulunmayan standart sütunlar: " + ", ".join(missing))
    report.event("INFO", "HEADER_FOUND", group.name, file, sheet, header.start,
                 f"Başlık: {header.start}; yükseklik: {header.height}; eşleşen: "
                 f"{sum(m.target is not None for m in header.matches)}")
    return mapping


def required_columns_reason(data: dict[str, Any], header: Header, mapping: dict[int, str],
                            cfg: dict[str, Any]) -> str | None:
    """Kaynak başlığı veya eşlenen adla zorunlu alanları dolgudan önce kontrol eder."""
    for name in cfg.get("ROW_REQUIRED_COLUMNS", ()):
        key = normalize(name)
        targets = [mapping[item.index] for item in header.matches
                   if key in {normalize(item.label or item.raw), normalize(item.raw),
                              normalize(mapping[item.index])}]
        if not any(usable_value(data.get(target)) for target in targets):
            return "ZORUNLU_ALAN_BOŞ: " + name
    return None


def rejection_reason(data: dict[str, Any], cfg: dict[str, Any], row_number: int) -> str | None:
    filled = [v for v in data.values() if usable_value(v)]
    if not filled:
        return "FORMÜL_SONUCU_YOK" if any(nonempty(v) for v in data.values()) else "BOŞ"
    row_id_column = cfg["ROW_ID_COLUMN"]
    row_id = row_number if row_id_column is None else data.get(row_id_column)
    # Sanal satır numarası alt bilgi etiketi taşıyamaz; bu durumda ilk dolu alanı tara.
    footer_probe = row_id if row_id_column is not None and usable_value(row_id) else filled[0]
    if normalize(str(footer_probe)) in {normalize(s) for s in cfg["FOOTER_LABELS"]}:
        return "ALT_BİLGİ/TOPLAM"
    if len(filled) < cfg["MIN_ROW_VALUES"]:
        return "AZ_DOLU_ALAN"
    if cfg["ROW_REQUIRE_ANY"] and not any(usable_value(data.get(k)) for k in cfg["ROW_REQUIRE_ANY"]):
        return "ANAHTAR_ALAN_YOK"
    if cfg["ROW_ID_REGEX"] and not re.fullmatch(cfg["ROW_ID_REGEX"], str(row_id)):
        return "KAYIT_NO_BİÇİMİ"
    return None


def fingerprint(data: dict[str, Any], keys: Sequence[str]) -> str:
    # Tür bilgisi dahil edilir: metin '001', sayı 1 ile eşit sayılmaz.
    keys = list(keys) or sorted(k for k, v in data.items() if nonempty(v))
    return json.dumps([(k, type(data.get(k)).__name__, data.get(k)) for k in keys],
                      ensure_ascii=False, default=str, separators=(",", ":"))


def quarantine_row(report: Report, group: str, file: str, sheet: str, row: SourceRow,
                   reason: str, data: dict[str, Any] | None = None) -> None:
    report.quarantine.append({"Klasör": group, "Dosya": file, "Sayfa": sheet,
                              "Satır": row.number, "Neden": reason,
                              "Ham Satır": json.dumps(row.values, ensure_ascii=False, default=str),
                              "Eşlenen Veri": json.dumps(data or {}, ensure_ascii=False, default=str)})


def extract_sheet(rows: Iterable[SourceRow], group: Group, report: Report, file: str,
                  sheet: str, matcher: HeaderMatcher, cfg: dict[str, Any]) -> bool:
    """Okuyucudan bağımsız dönüşüm çekirdeği. Her dolu veri satırı korunur veya karantinaya gider."""
    iterator = iter(rows)
    preview = list(itertools.islice(iterator, cfg["HEADER_SEARCH_END"] + max(cfg["HEADER_HEIGHTS"])))
    header = matcher.detect(preview)
    if not header:
        report.event("WARNING", "HEADER_NOT_FOUND", group.name, file, sheet,
                     detail=f"{cfg['HEADER_SEARCH_START']}–{cfg['HEADER_SEARCH_END']} aralığında "
                            "eşiklere/zorunlu alanlara uyan başlık yok; sayfa birleştirilmedi.")
        return False
    mapping = register_header(header, group, report, file, sheet, cfg, matcher.schema)
    group.table_count += 1
    group.read_files.add(file)
    source = itertools.chain((r for r in preview if r.number >= header.start + header.height), iterator)
    window = deque(itertools.islice(source, max(cfg["HEADER_HEIGHTS"])))
    previous: dict[str, Any] = {}
    before, rejected = len(group.records), 0
    while window:
        row = window[0]
        # Normal sayısal kayıtlar için pahalı başlık aramasını tekrar çalıştırmayın.
        id_column = cfg["ROW_ID_COLUMN"]
        id_index = next((i for i, name in mapping.items() if name == id_column), None)
        value = row.values[id_index] if id_index is not None and id_index < len(row.values) else None
        # Sanal ID her zaman sayısal olacağından None durumunda bu optimizasyonu kullanmak,
        # veri arasındaki tekrar başlıkların algılanmasını tamamen devre dışı bırakırdı.
        numeric_id = (id_column is not None and usable_value(value)
                      and re.fullmatch(r"[0-9]+(?:\.[0-9]+)?", str(value).strip()))
        repeated = [] if numeric_id else [matcher.candidate(list(window), 0, h) for h in cfg["HEADER_HEIGHTS"]]
        again = max((h for h in repeated if h), key=lambda h: h.quality, default=None)
        consumed = again.height if again else 1
        if again:
            header = again
            mapping = register_header(again, group, report, file, sheet, cfg, matcher.schema)
            previous.clear()
            report.event("INFO", "REPEATED_HEADER", group.name, file, sheet, row.number,
                         "Tekrar eden başlık veri sayılmadı; sütun eşleştirmesi yenilendi.")
        elif any(nonempty(v) for v in row.values):
            # Başlığı boş veya başlığın sağında olup veri içeren sütunlar da kaybolmaz.
            for index, val in enumerate(row.values):
                if nonempty(val) and index not in mapping:
                    raw = f"Başlıksız {excel_col(index)}"
                    mapping[index] = extra_name(group, raw, 1, cfg, matcher.schema)
                    report.event("WARNING", "UNNAMED_COLUMN", group.name, file, sheet, row.number,
                                 f"{excel_col(index)} sütunu -> {mapping[index]}")
            data = {target: source_text(row.values[i] if i < len(row.values) else None,
                                       row.formats[i] if i < len(row.formats) else "", cfg)
                    for i, target in mapping.items()}
            reason = required_columns_reason(data, header, mapping, cfg)
            # Dolgu bilinçli olarak opt-in. Boş satır/başlık/karantina sınırında sıfırlanır.
            for key in cfg["FILL_DOWN_COLUMNS"]:
                if not reason and not nonempty(data.get(key)) and key in previous:
                    data[key] = previous[key]
            for col, formula in row.missing_formulas:
                report.event("WARNING", "FORMULA_CACHE_MISSING", group.name, file, sheet, row.number,
                             f"{excel_col(col)}: Hesaplanmış değer yok; formül metin olarak korundu: {formula}")
            reason = reason or rejection_reason(data, cfg, row.number)
            if not reason and cfg["DROP_DUPLICATES"]:
                fp = fingerprint(data, cfg["DEDUP_KEYS"])
                if fp in group.seen:
                    reason = "TEKRAR_KAYIT"
                else:
                    group.seen.add(fp)
            if reason:
                quarantine_row(report, group.name, file, sheet, row, reason, data)
                previous.clear()
                rejected += 1
            else:
                group.records.append(Record(data, file, sheet, row.number))
                previous.update({k: v for k, v in data.items() if nonempty(v)})
        else:
            previous.clear()
        for _ in range(consumed):
            if window:
                window.popleft()
            nxt = next(source, None)
            if nxt is not None:
                window.append(nxt)
    if rejected:
        report.event("WARNING", "QUARANTINED_ROWS", group.name, file, sheet,
                     detail=f"{rejected} dolu satır Karantina'ya alındı; ana kayıtlara katılmadı.")
    report.event("INFO", "SHEET_DONE", group.name, file, sheet,
                 detail=f"Eklenen kayıt: {len(group.records) - before}; karantina: {rejected}")
    return True


def is_own_report(path: Path) -> bool:
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return False
    try:
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo("docProps/custom.xml")
            return info.file_size < 1_000_000 and REPORT_MARKER.encode() in archive.read(info)
    except (OSError, KeyError, zipfile.BadZipFile):
        return False


def report_group_mode(cfg: dict[str, Any]) -> str:
    """Eski harici config sözlüklerinde mevcut klasör davranışını korur."""
    return cfg.get("REPORT_GROUP_MODE", "subfolder")


def group_dimension_label(cfg: dict[str, Any]) -> str:
    return "Kaynak Sayfa" if report_group_mode(cfg) == "sheet_name" else "Klasör"


def sheet_name_group(report: Report, sheet: str, path: Path) -> Group:
    """Aynı adlı kaynak sayfaları büyük/küçük harf farkından bağımsız birleştirir."""
    group = next((item for item in report.groups if item.name.casefold() == sheet.casefold()), None)
    if group is None:
        group = Group(sheet)
        report.groups.append(group)
    if path not in group.files:
        group.files.append(path)
    return group


def discover(report: Report, cfg: dict[str, Any]) -> None:
    def acceptable(path: Path) -> bool:
        if path.suffix.lower() not in cfg["EXTENSIONS"]:
            return False
        if path.is_symlink() or path.resolve() == report.output:
            return False
        if any(fnmatch.fnmatchcase(path.name.casefold(), p.casefold()) for p in cfg["EXCLUDE_FILES"]):
            return False
        if is_own_report(path):
            report.event("INFO", "PREVIOUS_REPORT_SKIPPED", file=str(path.relative_to(report.root)),
                         detail="Bu programın önceki raporu kaynak veriye tekrar katılmadı.")
            return False
        return True

    def walk_error(exc: OSError) -> None:
        report.event("ERROR", "DIRECTORY_READ_FAILED", file=str(exc.filename), detail=str(exc))

    roots = sorted(report.root.iterdir(), key=lambda p: p.name.casefold())
    if report_group_mode(cfg) == "sheet_name":
        for folder, subdirs, filenames in os.walk(report.root, followlinks=False, onerror=walk_error):
            subdirs[:] = sorted((name for name in subdirs
                                 if name not in cfg["EXCLUDE_DIRS"]
                                 and not (Path(folder) / name).is_symlink()), key=str.casefold)
            report.source_files.extend(Path(folder) / name
                                       for name in sorted(filenames, key=str.casefold)
                                       if acceptable(Path(folder) / name))
        report.source_files.sort(key=lambda path: str(path.relative_to(report.root)).casefold())
        return

    for directory in roots:
        if not directory.is_dir() or directory.is_symlink() or directory.name in cfg["EXCLUDE_DIRS"]:
            continue
        group = Group(directory.name)
        for folder, subdirs, filenames in os.walk(directory, followlinks=False, onerror=walk_error):
            subdirs[:] = sorted((n for n in subdirs if n not in cfg["EXCLUDE_DIRS"]
                                 and not (Path(folder) / n).is_symlink()), key=str.casefold)
            group.files.extend(Path(folder) / n for n in sorted(filenames, key=str.casefold)
                               if acceptable(Path(folder) / n))
        # Boş ana klasör de bir sayfayı temsil eder; raporda görünür.
        report.groups.append(group)
    root_files = [p for p in roots if p.is_file() and acceptable(p)]
    if root_files and cfg["INCLUDE_ROOT_FILES"]:
        root_name = cfg["ROOT_GROUP_NAME"]
        existing = {g.name.casefold() for g in report.groups}
        while root_name.casefold() in existing:
            root_name += "_"
        report.groups.append(Group(root_name, files=root_files))
    elif root_files:
        for p in root_files:
            report.event("WARNING", "ROOT_FILE_SKIPPED", file=p.name,
                         detail="Dosya doğrudan kökte; INCLUDE_ROOT_FILES=False.")


def sheet_allowed(name: str, visible: bool, cfg: dict[str, Any]) -> bool:
    return ((visible or cfg["INCLUDE_HIDDEN_SHEETS"])
            and not any(fnmatch.fnmatchcase(name.casefold(), p.casefold()) for p in cfg["EXCLUDE_SHEETS"]))


def process_rows(rows: Iterable[SourceRow], group: Group, report: Report, file: str,
                 sheet: str, matcher: HeaderMatcher, cfg: dict[str, Any]) -> None:
    # Yarım okunan bir sayfadan kısmi kayıt bırakılmaz; dosyanın sağlam sayfaları korunur.
    nrecords, nquarantine = len(group.records), len(report.quarantine)
    tables, read_files = group.table_count, group.read_files.copy()
    seen = group.seen.copy() if cfg["DROP_DUPLICATES"] else None
    try:
        extract_sheet(rows, group, report, file, sheet, matcher, cfg)
    except Exception as exc:
        del group.records[nrecords:]
        del report.quarantine[nquarantine:]
        group.table_count, group.read_files = tables, read_files
        if seen is not None:
            group.seen = seen
        report.event("ERROR", "SHEET_READ_FAILED", group.name, file, sheet,
                     detail=f"Sayfanın kısmi aktarımı geri alındı: {type(exc).__name__}: {exc}")


def ooxml_rows(formula_sheet: Any, value_sheet: Any) -> Iterator[SourceRow]:
    # Yanlış kayıtlı worksheet dimension bilgisine güvenilmez.
    formula_sheet.reset_dimensions()
    value_sheet.reset_dimensions()
    streams = itertools.zip_longest(formula_sheet.iter_rows(), value_sheet.iter_rows(), fillvalue=())
    for number, (formulas, values) in enumerate(streams, 1):
        data, formats, missing = [], [], []
        for index, (fcell, vcell) in enumerate(itertools.zip_longest(formulas, values)):
            fval, vval = getattr(fcell, "value", None), getattr(vcell, "value", None)
            if (getattr(fcell, "data_type", None) == "f" and vval is None
                    and getattr(vcell, "data_type", None) not in ("str", "s", "inlineStr")):
                formula = fval if isinstance(fval, str) else getattr(fval, "text", str(fval))
                vval = f"{MISSING_FORMULA_PREFIX} {formula}"
                missing.append((index, str(formula)))
            data.append(vval)
            formats.append(getattr(fcell, "number_format", "General"))
        yield SourceRow(number, data, formats, missing)


def read_ooxml(path: Path, group: Group | None, report: Report, matcher: HeaderMatcher,
                cfg: dict[str, Any]) -> None:
    # İçe aktarımlar geciktirilir: çekirdek testleri Excel kütüphaneleri olmadan çalışır.
    from contextlib import ExitStack
    from openpyxl import load_workbook

    relative = str(path.relative_to(report.root))
    with ExitStack() as stack:
        fw = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        stack.callback(fw.close)
        vw = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        stack.callback(vw.close)
        for fs in fw.worksheets:
            if not sheet_allowed(fs.title, fs.sheet_state == "visible", cfg):
                report.event("INFO", "SHEET_EXCLUDED", group.name if group else fs.title,
                             relative, fs.title,
                             detail="Gizlilik/sayfa adı filtresi nedeniyle dışlandı.")
                continue
            target_group = group or sheet_name_group(report, fs.title, path)
            process_rows(ooxml_rows(fs, vw[fs.title]), target_group, report,
                         relative, fs.title, matcher, cfg)


def read_legacy(path: Path, group: Group | None, report: Report, matcher: HeaderMatcher,
                cfg: dict[str, Any]) -> None:
    from python_calamine import CalamineWorkbook, SheetTypeEnum, SheetVisibleEnum

    relative = str(path.relative_to(report.root))
    report.event("WARNING", "LEGACY_VALUE_MODE", group.name if group else "", relative,
                 detail="XLS/XLSB değer modunda okunuyor. Formül önbelleği eksikliği ve özel sayı "
                        "gösterimleri hücre bazında doğrulanamaz; kritik kaynakları XLSX olarak kaydedin.")
    with CalamineWorkbook.from_path(str(path)) as workbook:
        for metadata in workbook.sheets_metadata:
            if metadata.typ != SheetTypeEnum.WorkSheet:
                continue
            name = metadata.name
            if not sheet_allowed(name, metadata.visible == SheetVisibleEnum.Visible, cfg):
                report.event("INFO", "SHEET_EXCLUDED", group.name if group else name,
                             relative, name,
                             detail="Gizlilik/sayfa adı filtresi nedeniyle dışlandı.")
                continue
            target_group = group or sheet_name_group(report, name, path)
            # skip_empty_area=False kritik: 12. satırın gerçek Excel numarası korunur.
            def row_stream(sheet_name: str = name) -> Iterator[SourceRow]:
                raw = workbook.get_sheet_by_name(sheet_name).to_python(skip_empty_area=False)
                for n, values in enumerate(raw, 1):
                    yield SourceRow(n, list(values))
            process_rows(row_stream(), target_group, report, relative, name, matcher, cfg)


def collect(report: Report, cfg: dict[str, Any], schema: Sequence[ColumnSpec]) -> None:
    matcher = HeaderMatcher(schema, cfg)

    def collect_file(path: Path, group: Group | None) -> None:
        relative = str(path.relative_to(report.root))
        LOG.info("Okunuyor: %s", relative)
        try:
            reader = read_legacy if path.suffix.lower() in (".xls", ".xlsb") else read_ooxml
            reader(path, group, report, matcher, cfg)
        except Exception as exc:
            report.event("ERROR", "FILE_READ_FAILED", group.name if group else "", relative,
                         detail=f"{type(exc).__name__}: {exc}")

    if report_group_mode(cfg) == "sheet_name":
        for path in report.source_files:
            collect_file(path, None)
    else:
        for group in report.groups:
            if not group.files:
                report.event("WARNING", "EMPTY_GROUP", group.name,
                             detail="Bu ana klasörde kaynak Excel bulunamadı.")
            for path in group.files:
                collect_file(path, group)
    if not any(g.records for g in report.groups):
        report.event("WARNING", "NO_DATA", detail="Hiçbir veri kaydı birleştirilemedi; başlık eşiklerini/Denetim'i inceleyin.")


def category_text(value: Any) -> str:
    if not nonempty(value):
        return "(Boş)"
    if not usable_value(value):
        return "(Formül sonucu okunamadı)"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def category_counts(report: Report, dimension: str) -> Counter[str]:
    if dimension == "Klasör":
        return Counter({g.name: len(g.records) for g in report.groups if g.records})
    # Değerlerin noktalaması / baştaki sıfırı korunur; fuzzy SADECE başlıklardadır.
    return Counter(category_text(r.values.get(dimension)) for g in report.groups for r in g.records)


def top_categories(counts: Counter[str], limit: int) -> list[tuple[str, int]]:
    items = sorted(counts.items(), key=lambda x: (-x[1], x[0].casefold()))
    if len(items) <= limit:
        return items
    label = "Diğer kategoriler (toplam)"
    while label in counts:
        label += "*"
    return items[:limit] + [(label, sum(v for _, v in items[limit:]))]


def group_stats(report: Report, schema: Sequence[ColumnSpec]) -> list[dict[str, Any]]:
    output = []
    for group in report.groups:
        filled = sum(usable_value(r.values.get(s.name)) for r in group.records for s in schema)
        total_fields = len(group.records) * len(schema)
        warnings = sum(e["Klasör"] == group.name and e["Seviye"] in ("WARNING", "ERROR")
                       for e in report.events)
        quarantined = sum(q["Klasör"] == group.name for q in report.quarantine)
        output.append(dict(group=group, records=len(group.records), files=len(group.read_files),
                           found=len(group.files), sheets=group.table_count, filled=filled,
                           possible=total_fields, completeness=filled / total_fields if total_fields else 0,
                           warnings=warnings, quarantine=quarantined))
    return output


def workflow_counts(records: Iterable[Record], settings: dict[str, Any]) -> Counter[str]:
    """Durum önekini sayar; yalnızca açık kayıtlarda şirket hücresini inceler.

    Eksik şirket sütunu boş hücre sayılmaz. Formül sonucu okunamadığında da
    sıra tahmin edilmez. Sıfır ve False dolu hücredir; boşluk metni boş hücredir.
    """
    counts: Counter[str] = Counter()
    open_prefix = settings["OPEN_PREFIX"].strip().casefold()
    closed_prefix = settings["CLOSED_PREFIX"].strip().casefold()
    company_column = settings["COMPANY_COLUMN"]
    for record in records:
        status = record.values.get(settings["STATUS_COLUMN"])
        status = status.strip().casefold() if isinstance(status, str) and usable_value(status) else ""
        if status.startswith(closed_prefix):
            counts["closed"] += 1
            continue
        if not status.startswith(open_prefix):
            counts["unknown_status"] += 1
            continue
        counts["open"] += 1
        company = record.values.get(company_column)
        if company_column not in record.values or (nonempty(company) and not usable_value(company)):
            counts["unknown_turn"] += 1
        else:
            counts["other_company" if nonempty(company) else "our_company"] += 1
    return counts


def build_formats(workbook: Any, cfg: dict[str, Any]) -> dict[str, Any]:
    base = {"font_name": "Calibri", "font_size": 11, "font_color": PALETTE["ink"], "valign": "vcenter"}
    definitions = {
        "body": {"text_wrap": True},
        "title": {"font_size": 22, "bold": True, "font_color": "white", "bg_color": PALETTE["navy"]},
        "subtitle": {"font_size": 10, "font_color": PALETTE["muted"], "text_wrap": True},
        "section": {"font_size": 13, "bold": True, "font_color": PALETTE["navy"]},
        "header": {"bold": True, "font_color": "white", "bg_color": PALETTE["navy"], "text_wrap": True},
        "link": {"font_color": PALETTE["teal"], "underline": True},
        "integer": {"num_format": "#,##0"},
        "percent": {"num_format": "0.0%"},
        "date": {"num_format": cfg["DATE_FORMAT"]},
        "datetime": {"num_format": cfg["DATETIME_FORMAT"]},
        "time": {"num_format": "hh:mm:ss"},
        "duration": {"num_format": "[h]:mm:ss"},
        "warn": {"font_color": "#9A5705", "bg_color": "#FFF2D6", "text_wrap": True},
        "error": {"font_color": "#A32943", "bg_color": "#FFE6EC", "text_wrap": True},
        "ok": {"font_color": "#11735B", "bg_color": "#DCF7EF", "text_wrap": True},
        "pale": {"bg_color": PALETTE["paper"]},
    }
    return {name: workbook.add_format({**base, **options}) for name, options in definitions.items()}


def write_value(sheet: Any, row: int, col: int, value: Any, formats: dict[str, Any],
                default: str = "body") -> None:
    """Kaynak metinleri formül/URL olarak yorumlatmaz; uzun metni sessizce kesmez."""
    fmt = formats[default]
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            value = value.isoformat()  # Saat dilimini sessizce düşürmek yerine metin sakla.
        else:
            fmt = formats["datetime" if value.time() != time() else "date"]
    elif isinstance(value, date):
        fmt = formats["date"]
    elif isinstance(value, time):
        fmt = formats["time"]
    elif isinstance(value, timedelta):
        value, fmt = value.total_seconds() / 86400, formats["duration"]
    if isinstance(value, str):
        if len(value) > 32767:
            raise ValueError(f"Excel hücre metni sınırı aşıldı: {sheet.name}!{excel_col(col)}{row + 1}")
        result = sheet.write_string(row, col, value, fmt)
    else:
        result = sheet.write(row, col, value, fmt)
    if isinstance(result, int) and result < 0:
        raise ValueError(f"Hücre yazma hatası ({result}): {sheet.name}!{excel_col(col)}{row + 1}")


def internal_link(sheet_name: str, cell: str = "A1") -> str:
    return "internal:'" + sheet_name.replace("'", "''") + "'!" + cell


def write_group_sheets(workbook: Any, report: Report, cfg: dict[str, Any],
                       schema: Sequence[ColumnSpec], formats: dict[str, Any], used: set[str]) -> None:
    table_no = 0
    for group in report.groups:
        headers = [s.name for s in schema] + list(group.extra_names.values()) + list(META_COLUMNS)
        if len(headers) > 16384:
            raise ValueError(f"{group.name}: Excel sütun sınırı aşılıyor; sütun eşleştirmelerini daraltın.")
        group.sheet_names.clear()
        batches = range(0, max(1, len(group.records)), cfg["MAX_ROWS_PER_SHEET"])
        for part, start in enumerate(batches, 1):
            name = unique_sheet_name(group.name if part == 1 else f"{group.name} ({part})", used)
            group.sheet_names.append(name)
            ws = workbook.add_worksheet(name)
            ws.hide_gridlines(2)
            ws.set_tab_color(PALETTE["teal"])
            ws.set_zoom(90)
            ws.freeze_panes(7, 2)
            ws.set_default_row(32)
            ws.set_row(0, 28)
            ws.set_row(1, 28)
            ws.merge_range(0, 0, 1, len(headers) - 1, group.name, formats["title"])
            ws.write_url(3, 0, internal_link(cfg["SUMMARY_SHEET"]), formats["link"], "Özete dön")
            data = group.records[start:start + cfg["MAX_ROWS_PER_SHEET"]]
            ws.merge_range(4, 0, 4, len(headers) - 1,
                           f"Bölüm {part}  |  {len(data):,} kayıt  |  Kaynaklar: {len(group.read_files)} dosya / "
                           f"{group.table_count} sayfa  |  Satırlar kaynak sırasını korur.", formats["subtitle"])
            ws.set_row(4, 25)
            ws.set_row(6, 34)
            for col, header in enumerate(headers):
                width = next((s.width for s in schema if s.name == header), 36)
                if header == "Kaynak Dosya":
                    width = 52
                elif header == "Kaynak Satır":
                    width = 14
                ws.set_column(col, col, width)
            for r, record in enumerate(data, 7):
                values = [record.values.get(h) for h in headers[:-3]] + [record.file, record.sheet, record.row]
                for col, value in enumerate(values):
                    write_value(ws, r, col, value, formats)
            if data:
                table_no += 1
                ws.add_table(6, 0, 6 + len(data), len(headers) - 1,
                             {"name": f"ReportData_{table_no}", "style": "Table Style Medium 2",
                              "columns": [{"header": h, "header_format": formats["header"]} for h in headers]})
                ws.conditional_format(7, 0, 6 + len(data), len(schema) - 1,
                                      {"type": "blanks", "format": formats["warn"]})
                ws.conditional_format(7, 0, 6 + len(data), len(headers) - 1,
                                      {"type": "text", "criteria": "containing",
                                       "value": MISSING_FORMULA_PREFIX, "format": formats["error"]})
            else:
                for col, header in enumerate(headers):
                    write_value(ws, 6, col, header, formats, "header")
                ws.merge_range(7, 0, 7, len(headers) - 1,
                               "Birleştirilen kayıt yok. Ayrıntı için Denetim sayfasını inceleyin.", formats["warn"])
            ws.set_landscape()
            ws.fit_to_pages(1, 0)
            ws.repeat_rows(6)
            ws.set_footer("&L" + name + "&RSayfa &P / &N")
            ws.print_area(0, 0, max(7, 6 + len(data)), len(headers) - 1)


def write_audit_sheet(workbook: Any, title: str, rows: list[dict[str, Any]], headers: list[str],
                      formats: dict[str, Any], cfg: dict[str, Any], used: set[str]) -> None:
    # Uzun JSON satırlarını kesmek yerine ayrı hücre parçalarına böl.
    expanded_headers = []
    part_counts = {}
    for header in headers:
        parts = max((math.ceil(len(str(r.get(header, ""))) / 30000) for r in rows), default=1)
        part_counts[header] = max(1, parts)
        expanded_headers.extend(header if i == 0 else f"{header} (devam {i + 1})" for i in range(max(1, parts)))
    if len(expanded_headers) > 16384:
        raise ValueError(f"{title}: Denetim sütun sınırı aşıldı; JSONL dosyasına bakın.")
    for part, start in enumerate(range(0, max(1, len(rows)), cfg["MAX_ROWS_PER_SHEET"]), 1):
        name = title if part == 1 else unique_sheet_name(f"{title} ({part})", used)
        ws = workbook.add_worksheet(name)
        ws.hide_gridlines(2)
        ws.set_tab_color(PALETTE["amber"])
        ws.freeze_panes(4, 0)
        ws.set_default_row(32)
        ws.set_column(0, len(expanded_headers) - 1, 24)
        ws.merge_range(0, 0, 1, max(3, len(expanded_headers) - 1), title, formats["title"])
        ws.write_url(2, 0, internal_link(cfg["SUMMARY_SHEET"]), formats["link"], "Özete dön")
        segment = rows[start:start + cfg["MAX_ROWS_PER_SHEET"]]
        for r, record in enumerate(segment, 4):
            c = 0
            for header in headers:
                val = record.get(header)
                parts = part_counts[header]
                for i in range(parts):
                    value = str(val)[i * 30000:(i + 1) * 30000] if parts > 1 else val
                    write_value(ws, r, c, value, formats)
                    c += 1
        if segment:
            ws.add_table(3, 0, 3 + len(segment), len(expanded_headers) - 1,
                         {"style": "Table Style Medium 2",
                          "columns": [{"header": h, "header_format": formats["header"]} for h in expanded_headers]})
        else:
            for c, h in enumerate(expanded_headers):
                write_value(ws, 3, c, h, formats, "header")
            ws.write(4, 0, "Kayıt yok.", formats["ok"])
        if title == "Denetim" and segment:
            ws.conditional_format(4, 0, 3 + len(segment), len(expanded_headers) - 1,
                                  {"type": "formula", "criteria": '=$A5="ERROR"', "format": formats["error"]})
            ws.conditional_format(4, 0, 3 + len(segment), len(expanded_headers) - 1,
                                  {"type": "formula", "criteria": '=$A5="WARNING"', "format": formats["warn"]})
        for col, h in enumerate(expanded_headers):
            if h in ("Açıklama", "Ham Satır", "Eşlenen Veri", "Dosya") or "devam" in h:
                ws.set_column(col, col, 55)


def make_chart(workbook: Any, helper: Any, start_row: int, dimension: str, kind: str,
               categories: list[tuple[str, int]], color_index: int, formats: dict[str, Any]) -> Any:
    colors = [PALETTE[k] for k in ("teal", "cyan", "violet", "amber", "rose")]
    helper.write(start_row, 0, dimension, formats["header"])
    helper.write(start_row, 1, "Kayıt", formats["header"])
    for i, (label, count) in enumerate(categories, start_row + 1):
        write_value(helper, i, 0, label, formats)
        write_value(helper, i, 1, count, formats, "integer")
    chart = workbook.add_chart({"type": kind})
    series = {"name": "Kayıt sayısı", "categories": [helper.name, start_row + 1, 0, start_row + len(categories), 0],
              "values": [helper.name, start_row + 1, 1, start_row + len(categories), 1],
              "fill": {"color": colors[color_index % len(colors)]}, "border": {"none": True}}
    if kind == "doughnut":
        series["points"] = [{"fill": {"color": colors[i % len(colors)]}, "border": {"color": "white"}}
                            for i in range(len(categories))]
        series["data_labels"] = {"percentage": True, "font": {"size": 10}}
        chart.set_hole_size(65)
        chart.set_legend({"position": "right", "font": {"size": 9}})
    else:
        series["data_labels"] = {"value": True, "position": "outside_end", "font": {"size": 9}}
        chart.set_legend({"none": True})
        if kind == "bar":
            chart.set_x_axis({"num_format": "0", "min": 0, "major_gridlines": {"visible": True,
                              "line": {"color": PALETTE["line"]}}, "num_font": {"size": 9}})
            chart.set_y_axis({"reverse": True, "num_font": {"size": 9}})
        else:
            chart.set_y_axis({"num_format": "0", "min": 0, "major_gridlines": {"visible": True,
                              "line": {"color": PALETTE["line"]}}, "num_font": {"size": 9}})
            chart.set_x_axis({"num_font": {"size": 9, "rotation": -35}})
    chart.add_series(series)
    chart.set_title({"name": f"{dimension} · Kayıt dağılımı", "name_font": {"size": 13,
                     "bold": True, "color": PALETTE["navy"]}})
    chart.set_chartarea({"fill": {"color": "white"}, "border": {"color": PALETTE["line"]}})
    chart.set_plotarea({"fill": {"color": "white"}, "border": {"none": True}})
    chart.set_size({"width": 540, "height": 315})
    chart.show_hidden_data()
    return chart


def write_workflow_summary(ws: Any, report: Report, settings: dict[str, Any],
                           formats: dict[str, Any], group_label: str) -> None:
    """Genel ve rapor grubu bazındaki durum/sıra sayılarını dashboard'a yazar."""
    per_group = [(group.name, workflow_counts(group.records, settings)) for group in report.groups]
    totals: Counter[str] = Counter()
    for _, counts in per_group:
        totals.update(counts)
    ws.merge_range(11, 0, 11, 19, "AÇIK / KAPALI VE İŞ SIRASI", formats["section"])
    spans = [(0, 4), (5, 6), (7, 8), (9, 11), (12, 14), (15, 17), (18, 19)]
    labels = [group_label, "Açık", "Kapalı", "Sıra bizde", "Sıra diğer şirkette",
              "Durum belirsiz", "Sıra belirsiz (açık)"]
    keys = ("open", "closed", "our_company", "other_company", "unknown_status", "unknown_turn")
    for (first, last), label in zip(spans, labels):
        ws.merge_range(12, first, 12, last, label, formats["header"])
    ws.set_row(12, 48)
    for row, (label, counts) in enumerate([("GENEL TOPLAM", totals), *per_group], 13):
        values = [label, *(counts[key] for key in keys)]
        for (first, last), value in zip(spans, values):
            style = "body" if first == 0 else "integer"
            ws.merge_range(row, first, row, last, "", formats[style])
            write_value(ws, row, first, value, formats, style)
        ws.set_row(row, 28)
    note_row = 15 + len(per_group)
    ws.merge_range(note_row, 0, note_row + 1, 19,
                   f"Yalnızca açık kayıtlarda: {settings['COMPANY_COLUMN']} boşsa sıra bizde, doluysa diğer şirkette. "
                   "Kapalı kayıtlar sıra hesabına dahil değildir. Tanınmayan/eksik durumlar ve açık kayıtlardaki "
                   "eksik şirket sütunu veya okunamayan formül sonucu belirsiz olarak gösterilir.", formats["subtitle"])


def write_dashboard(workbook: Any, ws: Any, helper: Any, report: Report,
                    cfg: dict[str, Any], schema: Sequence[ColumnSpec], formats: dict[str, Any]) -> None:
    stats = group_stats(report, schema)
    group_label = group_dimension_label(cfg)
    processed_files = len({file for group in report.groups for file in group.read_files})
    found_files = (len(report.source_files) if report_group_mode(cfg) == "sheet_name"
                   else sum(item["found"] for item in stats))
    grouping_text = ("Kaynak Excel sayfası → rapor sayfası"
                     if report_group_mode(cfg) == "sheet_name"
                     else "İlk düzey klasör → rapor sayfası")
    total = sum(s["records"] for s in stats)
    errors = sum(e["Seviye"] == "ERROR" for e in report.events)
    warnings = sum(e["Seviye"] == "WARNING" for e in report.events)
    filled, possible = sum(s["filled"] for s in stats), sum(s["possible"] for s in stats)
    workflow = cfg.get("WORKFLOW_SUMMARY", {"ENABLED": False})
    workflow_offset = len(stats) + 6 if workflow["ENABLED"] else 0
    ws.hide_gridlines(2)
    ws.set_tab_color(PALETTE["navy"])
    ws.set_zoom(85)
    ws.set_column(0, 19, 7.8)
    ws.set_default_row(20)
    for r in range(0, 58 + len(stats) + workflow_offset):
        ws.set_row(r, 20, formats["pale"])
    ws.set_row(0, 29)
    ws.set_row(1, 29)
    ws.merge_range("A1:T2", cfg["REPORT_TITLE"], formats["title"])
    ws.merge_range("A3:T3", f"Üretim: {report.generated_at:%d.%m.%Y %H:%M}  |  "
                   f"{grouping_text}  |  Grafikte en yoğun kategoriler ve kalanların toplamı",
                   formats["subtitle"])
    # Yardımcı giriş verileri: Python'un oluşturduğu rapor anı istatistikleri.
    group_metric_label = "Kaynak sayfa" if report_group_mode(cfg) == "sheet_name" else "Ana klasör"
    metrics = [("Kayıt", total), ("İşlenen dosya", processed_files),
               (group_metric_label, len(stats)), ("Uyarı + hata", warnings + errors),
               ("Dolu standart alan", filled), ("Olası standart alan", possible),
               ("Karantina", len(report.quarantine))]
    helper.write_row(0, 3, ["Gösterge", "Rapor anındaki değer"], formats["header"])
    for i, pair in enumerate(metrics, 1):
        helper.write_row(i, 3, pair)
    for i, (label, value) in enumerate(metrics[:4]):
        c1, c2 = i * 5, i * 5 + 3
        color = [PALETTE[k] for k in ("teal", "cyan", "violet", "amber")][i]
        label_fmt = workbook.add_format({"font_name": "Calibri", "font_size": 10, "bold": True,
                                         "bg_color": color, "font_color": "white", "valign": "vcenter"})
        value_fmt = workbook.add_format({"font_name": "Calibri", "font_size": 28, "bold": True,
                                         "bg_color": "white", "font_color": color, "valign": "vcenter",
                                         "num_format": "#,##0"})
        ws.merge_range(4, c1, 4, c2, label.upper(), label_fmt)
        ws.merge_range(5, c1, 7, c2, "", value_fmt)
        ws.write_formula(5, c1, f"='_Rapor Verisi'!E{i + 2}", value_fmt, value)
    if errors:
        message, style = f"EKSİK RAPOR: {errors} okuma/erişim hatası var. Denetim sayfasındaki ERROR kayıtlarını inceleyin.", "error"
    elif warnings or report.quarantine:
        message, style = f"KONTROL GEREKİYOR: {warnings} uyarı; {len(report.quarantine)} dolu satır karantinada. Denetim ve eşleştirmeleri gözden geçirin.", "warn"
    else:
        message, style = "Aktarım tamamlandı. Eşleştirme kararları ve kaynak konumları denetim sayfalarında kayıtlıdır.", "ok"
    ws.merge_range("A10:T10", message, formats[style])
    ws.set_row(9, 29)
    if workflow["ENABLED"]:
        write_workflow_summary(ws, report, workflow, formats, group_label)
    positions = [(12, 0), (12, 10), (26, 0), (26, 10)]
    for i, (dimension, kind) in enumerate(cfg["CHARTS"]):
        categories = top_categories(category_counts(report, dimension), cfg["DASHBOARD_TOP_N"])
        chart_label = group_label if dimension == "Klasör" else dimension
        r, c = positions[i]
        r += workflow_offset
        if categories:
            chart = make_chart(workbook, helper, 12 + i * (cfg["DASHBOARD_TOP_N"] + 5),
                               chart_label, kind, categories, i, formats)
            ws.insert_chart(r, c, chart, {"x_offset": 3, "y_offset": 3})
        else:
            ws.merge_range(r, c, r + 11, c + 8,
                           f"{chart_label}\nGrafik için veri bulunamadı.", formats["subtitle"])
    ws.merge_range(40 + workflow_offset, 0, 40 + workflow_offset, 19,
                   f"{group_label.upper()} BAZINDA ÖZET", formats["section"])
    spans = [(0, 4), (5, 6), (7, 8), (9, 10), (11, 13), (14, 16), (17, 19)]
    first_label = "Kaynak Sayfa / Rapor Sayfası" if report_group_mode(cfg) == "sheet_name" else "Klasör / Sayfa"
    labels = [first_label, "Dosya", "Sayfa", "Kayıt", "Alan doluluğu", "Uyarı + hata", "Karantina"]
    for (c1, c2), label in zip(spans, labels):
        ws.merge_range(42 + workflow_offset, c1, 42 + workflow_offset, c2, label, formats["header"])
    ws.set_row(42 + workflow_offset, 32)
    for r, item in enumerate(stats, 43 + workflow_offset):
        g = item["group"]
        values = [g.name, item["files"], item["sheets"], item["records"], item["completeness"],
                  item["warnings"], item["quarantine"]]
        for idx, ((c1, c2), value) in enumerate(zip(spans, values)):
            fmt = formats["percent" if idx == 4 else "integer" if idx else "body"]
            ws.merge_range(r, c1, r, c2, "", fmt)
            write_value(ws, r, c1, value, formats, "percent" if idx == 4 else "integer" if idx else "body")
        ws.set_row(r, 28)
        if g.sheet_names:
            ws.write_url(r, 0, internal_link(g.sheet_names[0]), formats["link"], g.name)
    end = 44 + len(stats) + workflow_offset
    completeness = filled / possible if possible else 0
    ws.merge_range(end, 0, end, 19,
                   f"Genel alan doluluğu: {completeness:.1%}  |  Taranan dosya: {found_files}  |  "
                   f"Karantina: {len(report.quarantine)}  |  Hata: {errors}", formats["section"])
    ws.merge_range(end + 2, 0, end + 3, 19,
                   "Alan doluluğu = dolu standart hücre / (kayıt × tanımlı standart sütun). Bu değer doğruluk puanı değildir. "
                   "Dashboard rapor üretim anını gösterir; kaynak veya çıktı verileri değiştirilirse güncel rapor için kodu yeniden çalıştırın. "
                   "Karantina kayıtları ana sayfalara ve grafik toplamlarına dahil değildir.", formats["subtitle"])
    for col, title in ((0, "Denetim"), (6, "Başlık Eşleştirmeleri"), (14, "Karantina")):
        ws.write_url(end + 5, col, internal_link(title), formats["link"], title)
    ws.set_landscape()
    ws.set_paper(9)
    ws.fit_to_pages(1, 0)
    ws.set_margins(0.25, 0.25, 0.3, 0.3)
    ws.print_area(0, 0, end + 5, 19)
    ws.set_footer("&LExcel birleştirme raporu&RSayfa &P / &N")
    ws.activate()
    ws.set_first_sheet()


def render_report(workbook: Any, report: Report, cfg: dict[str, Any], schema: Sequence[ColumnSpec]) -> None:
    formats = build_formats(workbook, cfg)
    used = {n.casefold() for n in AUX_SHEETS}
    summary = unique_sheet_name(cfg["SUMMARY_SHEET"], used)
    # Hiperlinklerde sanitizasyondan sonraki gerçek adı kullan.
    local_cfg = {**cfg, "SUMMARY_SHEET": summary}
    ws = workbook.add_worksheet(summary)  # MUTLAKA İLK SAYFA.
    subject = ("Kaynak sayfa bazlı konsolide rapor" if report_group_mode(cfg) == "sheet_name"
               else "Klasör bazlı konsolide rapor")
    workbook.set_properties({"title": cfg["REPORT_TITLE"], "subject": subject,
                             "author": "Excel Raporlama", "comments": "Kaynak dosyalar değiştirilmemiştir."})
    workbook.set_custom_property("ReportGenerator", REPORT_MARKER)
    write_group_sheets(workbook, report, local_cfg, schema, formats, used)
    write_audit_sheet(workbook, "Denetim", report.events,
                      ["Seviye", "Kod", "Klasör", "Dosya", "Sayfa", "Satır", "Açıklama"], formats, local_cfg, used)
    write_audit_sheet(workbook, "Başlık Eşleştirmeleri", report.mappings,
                      ["Klasör", "Dosya", "Sayfa", "Başlık Satırı", "Yükseklik", "Kaynak Sütun", "Orijinal Başlık",
                       "Ayıklanan Başlık", "Başlık İşleme", "Skorlanan Başlık", "Hedef Sütun", "Eşleşme", "Skor", "İkinci Aday",
                       "Skor Farkı"], formats, local_cfg, used)
    write_audit_sheet(workbook, "Karantina", report.quarantine,
                      ["Klasör", "Dosya", "Sayfa", "Satır", "Neden", "Ham Satır", "Eşlenen Veri"], formats, local_cfg, used)
    helper = workbook.add_worksheet("_Rapor Verisi")
    write_dashboard(workbook, ws, helper, report, local_cfg, schema, formats)
    helper.hide()


def save_report(report: Report, cfg: dict[str, Any], schema: Sequence[ColumnSpec]) -> None:
    import xlsxwriter

    fd, temp_name = tempfile.mkstemp(prefix=".excel_rapor_", suffix=".xlsx", dir=report.output.parent)
    os.close(fd)
    temporary = Path(temp_name)
    try:
        # constant_memory kullanmayın: dashboard merge/table özellikleriyle uyumlu değil.
        with xlsxwriter.Workbook(str(temporary), {"strings_to_formulas": False,
                                  "strings_to_urls": False, "strings_to_numbers": False}) as workbook:
            render_report(workbook, report, cfg, schema)
        if report.output.exists() and not cfg["OVERWRITE"]:
            raise FileExistsError("Çıktı mevcut. Üzerine yazmak için --overwrite kullanın.")
        # Yarım oluşturulmuş rapor nihai dosyanın üstüne yazılmaz.
        os.replace(temporary, report.output)
    finally:
        temporary.unlink(missing_ok=True)


def save_json_audit(report: Report) -> Path:
    path = report.output.with_suffix(".denetim.jsonl")
    fd, name = tempfile.mkstemp(prefix=".denetim_", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            for section, records in (("event", report.events), ("header", report.mappings),
                                     ("quarantine", report.quarantine)):
                for record in records:
                    stream.write(json.dumps({"type": section, **record}, ensure_ascii=False, default=str) + "\n")
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)
    return path


def absolute_path(value: str | Path) -> Path:
    path = Path(value).expanduser()
    return (path if path.is_absolute() else SCRIPT_DIR / path).resolve()


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", help="Kaynak ana klasör; verilmezse CONFIG kullanılır.")
    parser.add_argument("--output", help="Çıktı XLSX yolu; verilmezse CONFIG kullanılır.")
    parser.add_argument("--overwrite", action="store_true", help="Mevcut çıktı raporunun üzerine yaz.")
    parser.add_argument("--dry-run", action="store_true", help="Kaynakları incele; Excel üretmeden JSONL denetimi çıkar.")
    args = parser.parse_args(argv)
    cfg = {**CONFIG, "OVERWRITE": args.overwrite or CONFIG["OVERWRITE"]}
    try:
        validate(cfg, COLUMN_SCHEMA)
        root = absolute_path(args.input or cfg["INPUT_DIR"])
        output = absolute_path(args.output or cfg["OUTPUT_FILE"])
        if not root.is_dir():
            raise NotADirectoryError(f"Kaynak klasör bulunamadı: {root}")
        if output.suffix.lower() != ".xlsx":
            raise ValueError("Çıktı uzantısı .xlsx olmalıdır.")
        if output.exists() and not cfg["OVERWRITE"] and not args.dry_run:
            raise FileExistsError(f"Çıktı mevcut: {output}. --overwrite kullanın.")
        if output.exists() and not is_own_report(output) and cfg["OVERWRITE"]:
            raise ValueError("Güvenlik: Var olan çıktı bu programın raporu değil; farklı çıktı adı seçin.")
        output.parent.mkdir(parents=True, exist_ok=True)
        logging.basicConfig(level=logging.INFO, format="%(levelname)s | %(message)s", force=True,
                            handlers=[logging.StreamHandler(),
                                      logging.FileHandler(output.with_suffix(".log"), encoding="utf-8", mode="w")])
        report = Report(root, output)
        discover(report, cfg)
        collect(report, cfg, COLUMN_SCHEMA)
        audit = save_json_audit(report)
        errors = sum(e["Seviye"] == "ERROR" for e in report.events)
        if cfg["FAIL_ON_FILE_ERROR"] and errors:
            raise RuntimeError(f"{errors} kaynak hatası nedeniyle çıktı üretilmedi. Denetim: {audit}")
        if not args.dry_run:
            save_report(report, cfg, COLUMN_SCHEMA)
        LOG.info("%s | %s kayıt | %s %s | %s hata | JSONL: %s",
                 "Ön kontrol tamamlandı" if args.dry_run else f"Rapor hazır: {output}",
                 sum(len(g.records) for g in report.groups), len(report.groups),
                 group_dimension_label(cfg).casefold(), errors, audit)
        return 2 if errors else 0
    except KeyboardInterrupt:
        print("İşlem kullanıcı tarafından durduruldu.", file=sys.stderr)
        return 130
    except Exception as exc:
        print(f"HATA: {type(exc).__name__}: {exc}", file=sys.stderr)
        print("Çıktı Excel'de açıksa kapatın. Ayrıntı için .log / .denetim.jsonl dosyasına bakın.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
