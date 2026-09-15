"""Integration regressions for the desktop adapter (no display required)."""
import copy
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

import openpyxl

import excel_fusion as core
from fusion_service import RunOptions, run_report, validate_paths


class ServiceTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.source = self.root / "kaynaklar" / "Grup A" / "Türkçe kaynak.xlsx"
        self.source.parent.mkdir(parents=True)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append([column.name for column in core.COLUMN_SCHEMA])
        sheet.append([1, "E1", "P1", "001", "Test", "open", None])
        workbook.save(self.source)
        workbook.close()
        self.options = RunOptions(str(self.source.parent.parent), str(self.root / "çıktı" / "Rapor.xlsx"))

    def test_report_preserves_source_and_global_config(self):
        before = self.source.read_bytes()
        original = copy.deepcopy(core.CONFIG)
        stages = []
        result = run_report(self.options, stages.append)
        self.assertEqual((result.records, result.groups, result.errors), (1, 1, 0))
        self.assertTrue(result.audit.is_file())
        workbook = openpyxl.load_workbook(result.output, data_only=True)
        try:
            self.assertIn("Grup A", workbook.sheetnames)
            self.assertEqual(workbook["Grup A"]["A8"].value, 1)
        finally:
            workbook.close()
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(core.CONFIG, original)
        self.assertGreaterEqual(len(stages), 3)

    def test_dry_run_only_writes_audit_and_log(self):
        result = run_report(replace(self.options, dry_run=True))
        self.assertIsNone(result.output)
        self.assertFalse(Path(self.options.output).exists())
        self.assertTrue(result.audit.is_file())
        self.assertTrue(Path(self.options.output).with_suffix(".log").is_file())

    def test_existing_output_protected_and_own_report_can_be_renewed(self):
        result = run_report(self.options)
        original = result.output.read_bytes()
        with self.assertRaises(FileExistsError):
            run_report(self.options)
        self.assertEqual(result.output.read_bytes(), original)
        renewed = run_report(replace(self.options, overwrite=True))
        self.assertEqual(renewed.records, 1)

    def test_source_cannot_be_overwritten(self):
        before = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, "raporu değil"):
            run_report(replace(self.options, output=str(self.source), overwrite=True))
        self.assertEqual(self.source.read_bytes(), before)

    def test_invalid_paths(self):
        for changes in ({"source": ""}, {"output": ""}, {"source": str(self.root / "missing")},
                        {"output": str(self.root / "bad.csv")}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_paths(replace(self.options, **changes))

    def test_invalid_settings(self):
        for changes in ({"header_row": 0}, {"search_end": 5}, {"search_end": 1048577},
                        {"group_mode": "invalid"}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                replace(self.options, **changes).config()

    def test_root_inclusion_and_sheet_mode(self):
        self.source.rename(self.source.parent.parent / "root.xlsx")
        for mode, include_root, expected in (("subfolder", False, 0), ("subfolder", True, 1),
                                              ("sheet_name", False, 1)):
            with self.subTest(mode=mode, include_root=include_root):
                result = run_report(replace(self.options, group_mode=mode,
                                            include_root=include_root, dry_run=True))
                self.assertEqual(result.records, expected)

    def test_partial_failure_reported_and_strict_mode_blocks_output(self):
        (self.source.parent / "broken.xlsx").write_bytes(b"invalid workbook")
        partial = run_report(self.options)
        self.assertEqual((partial.records, partial.errors), (1, 1))
        strict = replace(self.options, output=str(self.root / "strict.xlsx"), fail_on_error=True)
        with self.assertRaisesRegex(RuntimeError, "1 kaynak hatası"):
            run_report(strict)
        self.assertFalse(Path(strict.output).exists())
        self.assertTrue(Path(strict.output).with_suffix(".denetim.jsonl").exists())

    def test_handler_cleanup_on_save_failure(self):
        handlers, level = list(core.LOG.handlers), core.LOG.level
        with (patch.object(core, "save_report", side_effect=PermissionError("locked")),
              self.assertRaises(PermissionError)):
            run_report(self.options)
        self.assertEqual(core.LOG.handlers, handlers)
        self.assertEqual(core.LOG.level, level)
        self.assertEqual(run_report(self.options).records, 1)

    def test_adapter_and_cli_have_equivalent_report_data(self):
        cli_output = self.root / "cli.xlsx"
        with patch("sys.stderr"):
            self.assertEqual(core.main(["--input", self.options.source, "--output", str(cli_output)]), 0)
        # main configures global logging; close its file handler before temp cleanup on Windows.
        for handler in list(core.logging.getLogger().handlers):
            core.logging.getLogger().removeHandler(handler)
            handler.close()
        result = run_report(self.options)
        cli = openpyxl.load_workbook(cli_output, data_only=True)
        gui = openpyxl.load_workbook(result.output, data_only=True)
        try:
            self.assertEqual(cli.sheetnames, gui.sheetnames)
            self.assertEqual(list(cli["Grup A"].values), list(gui["Grup A"].values))
        finally:
            cli.close()
            gui.close()


if __name__ == "__main__":
    unittest.main()
