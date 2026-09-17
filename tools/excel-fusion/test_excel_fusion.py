"""Çekirdek ve XLSX içe aktarma/özet regresyon testleri.

Çalıştırma: python -m unittest discover -v
XLSX testleri uygulamanın openpyxl ve xlsxwriter bağımlılıklarını kullanır.
"""
import copy
import tempfile
import unittest
from pathlib import Path

import excel_fusion as fusion
from fusion_diagnostics import sample_configuration


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(sample_configuration())
        self.cfg = copy.deepcopy(fusion.CONFIG)
        self.settings = self.cfg["WORKFLOW_SUMMARY"]

    def count(self, values):
        records = [fusion.Record(value, "source.xlsx", "Data", row)
                   for row, value in enumerate(values, 1)]
        return fusion.workflow_counts(records, self.settings)

    def test_status_prefix_and_company_emptiness(self):
        for status in ("open", "  OpEn uzun açıklama\nclosed daha sonra", "OPEN: açıklama", "openDevam"):
            for company in (None, "", " \t\n", "yanıt", 0, False):
                with self.subTest(status=status, company=company):
                    counts = self.count([{"Open / Closed": status, "My Company": company}])
                    turn = "our_company" if company in (None, "", " \t\n") else "other_company"
                    self.assertEqual(counts, {"open": 1, turn: 1})

    def test_closed_never_counts_a_turn(self):
        rows = [{"Open / Closed": "  CLOSED uzun metin open"}]
        rows.extend({"Open / Closed": "closed", "My Company": value}
                    for value in (None, "yanıt", fusion.MISSING_FORMULA_PREFIX + " =A1"))
        self.assertEqual(self.count(rows), {"closed": 4})

    def test_unknown_status_and_unavailable_company_are_separate(self):
        rows = [{"Open / Closed": value} for value in
                (None, "", "pending open", "reopened", 1, fusion.MISSING_FORMULA_PREFIX + " =A1")]
        rows.extend([{}, {"Open / Closed": "open"},
                     {"Open / Closed": "open", "My Company": fusion.MISSING_FORMULA_PREFIX + " =A1"}])
        self.assertEqual(self.count(rows), {"unknown_status": 7, "open": 2, "unknown_turn": 2})

    def test_empty_records(self):
        self.assertEqual(self.count([]), {})

    def test_custom_columns_and_prefixes(self):
        self.settings.update(STATUS_COLUMN="Entity", COMPANY_COLUMN="Panel",
                             OPEN_PREFIX=" Aktif ", CLOSED_PREFIX=" Bitti ")
        fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)
        self.assertEqual(self.count([{"Entity": "AKTIF", "Panel": None},
                                     {"Entity": "aktif: bekliyor", "Panel": "cevap"},
                                     {"Entity": "bitti açıklama"}]),
                         {"open": 2, "our_company": 1, "other_company": 1, "closed": 1})

    def test_invalid_config(self):
        for overrides in ({"ENABLED": "yes"}, {"STATUS_COLUMN": "missing"},
                          {"COMPANY_COLUMN": "missing"}, {"COMPANY_COLUMN": "Open / Closed"},
                          {"OPEN_PREFIX": " "}, {"CLOSED_PREFIX": None},
                          {"OPEN_PREFIX": "closed"}, {"OPEN_PREFIX": "CLO"},
                          {"CLOSED_PREFIX": "op"}):
            with self.subTest(overrides=overrides):
                cfg = copy.deepcopy(self.cfg)
                cfg["WORKFLOW_SUMMARY"].update(overrides)
                with self.assertRaises(ValueError):
                    fusion.validate(cfg, fusion.COLUMN_SCHEMA)

    def test_invalid_report_group_mode(self):
        self.cfg["REPORT_GROUP_MODE"] = "folder_name"
        with self.assertRaisesRegex(ValueError, "REPORT_GROUP_MODE"):
            fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)

    def test_missing_report_group_mode_keeps_legacy_default(self):
        del self.cfg["REPORT_GROUP_MODE"]
        fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)
        self.assertEqual(fusion.report_group_mode(self.cfg), "subfolder")

    def test_disabled_and_legacy_config(self):
        self.cfg["WORKFLOW_SUMMARY"] = {"ENABLED": False}
        fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)
        del self.cfg["WORKFLOW_SUMMARY"]
        fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)

    def test_header_aliases(self):
        matcher = fusion.HeaderMatcher(fusion.COLUMN_SCHEMA, self.cfg)
        matches = matcher.match(["Open/Closed\nAçıklama", "My Company\nYanıtınızı yazın"])
        self.assertEqual([match.target for match in matches], ["Open / Closed", "My Company"])

    def test_none_row_id_column_uses_source_row_number_for_regex(self):
        self.cfg["ROW_ID_COLUMN"] = None
        self.cfg["ROW_ID_REGEX"] = r"^12$"

        fusion.validate(self.cfg, fusion.COLUMN_SCHEMA)

        data = {"No.": "sütundaki-değer", "Entity": "E"}
        self.assertIsNone(fusion.rejection_reason(data, self.cfg, 12))
        self.assertEqual(fusion.rejection_reason(data, self.cfg, 13), "KAYIT_NO_BİÇİMİ")

    def test_none_row_id_column_keeps_footer_detection(self):
        self.cfg["ROW_ID_COLUMN"] = None
        data = {"No.": "total", "Entity": None}

        self.assertEqual(fusion.rejection_reason(data, self.cfg, 12), "ALT_BİLGİ/TOPLAM")

    def test_none_row_id_column_keeps_repeated_header_detection(self):
        self.cfg["ROW_ID_COLUMN"] = None
        self.cfg["EXPECTED_HEADER_ROW"] = 1
        rows = [
            fusion.SourceRow(1, ["No.", "Entity", "Panel"]),
            fusion.SourceRow(2, ["A", "Entity A", "Panel A"]),
            fusion.SourceRow(3, ["No.", "Entity", "Panel"]),
            fusion.SourceRow(4, ["B", "Entity B", "Panel B"]),
        ]
        group = fusion.Group("Test")
        report = fusion.Report(Path("."), Path("report.xlsx"))
        matcher = fusion.HeaderMatcher(fusion.COLUMN_SCHEMA, self.cfg)

        self.assertTrue(fusion.extract_sheet(rows, group, report, "source.xlsx", "Data",
                                             matcher, self.cfg))
        self.assertEqual([record.values["No."] for record in group.records], ["A", "B"])
        self.assertTrue(any(event["Kod"] == "REPEATED_HEADER" for event in report.events))


class QuarantineTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(sample_configuration())

    def extract(self, values, **overrides):
        cfg = copy.deepcopy(fusion.CONFIG)
        cfg.update(EXPECTED_HEADER_ROW=1, HEADER_HEIGHTS=(1,), MIN_ROW_VALUES=1,
                   ROW_REQUIRE_ANY=())
        cfg.update(overrides)
        # Özel şema, kontrolün global COLUMN_SCHEMA'ya bağlı olmadığını doğrular.
        schema = [*fusion.COLUMN_SCHEMA, fusion.ColumnSpec("Custom")]
        rows = [fusion.SourceRow(1, ["No.", "Entity", "Panel", "Custom", "Notes"])]
        rows.extend(fusion.SourceRow(number, value) for number, value in enumerate(values, 2))
        group = fusion.Group("Test")
        report = fusion.Report(Path("."), Path("report.xlsx"))
        matcher = fusion.HeaderMatcher(schema, cfg)
        self.assertTrue(fusion.extract_sheet(rows, group, report, "source.xlsx", "Data",
                                            matcher, cfg))
        return group, report

    def test_required_empty_values_are_quarantined(self):
        for value in (None, "", " \t\n", fusion.MISSING_FORMULA_PREFIX + " =A1"):
            with self.subTest(value=value):
                group, report = self.extract([[value, None, None, None, "note", "unnamed"]],
                                             ROW_REQUIRED_COLUMNS=("No.",))
                self.assertEqual(group.records, [])
                self.assertEqual(len(report.quarantine), 1)
                rejected = report.quarantine[0]
                self.assertEqual(rejected["Neden"], "ZORUNLU_ALAN_BOŞ: No.")
                self.assertEqual(rejected["Satır"], 2)
                self.assertIn("unnamed", rejected["Ham Satır"])
                self.assertIn("note", rejected["Eşlenen Veri"])

    def test_required_values_include_zero_and_false(self):
        for value in ("text", 0, False):
            with self.subTest(value=value):
                group, report = self.extract([[None, None, None, value]],
                                             ROW_REQUIRED_COLUMNS=("Custom",))
                self.assertEqual(len(group.records), 1)
                self.assertEqual(group.records[0].values["Custom"], value)
                self.assertEqual(report.quarantine, [])

    def test_fill_down_cannot_rescue_required_cell(self):
        group, report = self.extract([
            [1, "E", "P"],
            [2, None, "P"],
            [3, "F", "P"],
        ], FILL_DOWN_COLUMNS=("Entity",), ROW_REQUIRED_COLUMNS=("Entity",))
        self.assertEqual([record.row for record in group.records], [2, 4])
        self.assertEqual(report.quarantine[0]["Neden"], "ZORUNLU_ALAN_BOŞ: Entity")

    def test_required_source_column_outside_schema_and_every_column_required(self):
        group, report = self.extract([
            [1, "E", "P", None, "note"],
            [2, "E", "P"],
            [None, "E", "P", None, "note"],
        ], ROW_REQUIRED_COLUMNS=("No.", " notes "))
        self.assertEqual(len(group.records), 1)
        self.assertEqual(len(report.quarantine), 2)

    def test_missing_header_is_rejected_before_other_rules(self):
        group, report = self.extract([["total", "E", "P"]],
                                     ROW_REQUIRED_COLUMNS=("Missing",))
        self.assertEqual(group.records, [])
        self.assertEqual(report.quarantine[0]["Neden"], "ZORUNLU_ALAN_BOŞ: Missing")

    def test_disabled_rule_allows_extra_only_row(self):
        group, report = self.extract([[None, None, None, None, "note"]])
        self.assertEqual(len(group.records), 1)
        self.assertEqual(report.quarantine, [])

    def test_required_columns_config_validation_and_legacy_default(self):
        cfg = copy.deepcopy(fusion.CONFIG)
        for value in (None, "Entity", [""], [" "], [1]):
            with self.subTest(value=value):
                cfg["ROW_REQUIRED_COLUMNS"] = value
                with self.assertRaisesRegex(ValueError, "ROW_REQUIRED_COLUMNS"):
                    fusion.validate(cfg, fusion.COLUMN_SCHEMA)
        cfg["ROW_REQUIRED_COLUMNS"] = ["Outside schema"]
        fusion.validate(cfg, fusion.COLUMN_SCHEMA)
        del cfg["ROW_REQUIRED_COLUMNS"]
        fusion.validate(cfg, fusion.COLUMN_SCHEMA)

    def test_existing_rules_and_empty_row_behavior_are_preserved(self):
        group, report = self.extract([
            [None, " ", None],
            [1],
            [2, "E"],
        ], MIN_ROW_VALUES=2)
        self.assertEqual(len(group.records), 1)
        self.assertEqual([row["Neden"] for row in report.quarantine], ["AZ_DOLU_ALAN"])


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.enterContext(sample_configuration())

    @staticmethod
    def create_source(path, sheets):
        import openpyxl

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for title, number in sheets:
            sheet = workbook.create_sheet(title)
            sheet.append(["No.", "Entity", "Panel"])
            sheet.append([number, f"Entity {number}", f"Panel {number}"])
        workbook.save(path)
        workbook.close()

    def test_sheet_name_mode_groups_same_sheets_across_all_directories(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "kaynaklar"
            self.create_source(root / "Kök.xlsx", [("Data", 1)])
            self.create_source(root / "Grup A" / "a.xlsx", [("Data", 2), ("Diğer", 20)])
            self.create_source(root / "Grup B" / "b.xlsx", [("data", 3)])

            cfg = copy.deepcopy(fusion.CONFIG)
            cfg["REPORT_GROUP_MODE"] = "sheet_name"
            cfg["INCLUDE_ROOT_FILES"] = False
            fusion.validate(cfg, fusion.COLUMN_SCHEMA)
            report = fusion.Report(root, root / "report.xlsx")

            fusion.discover(report, cfg)
            self.assertEqual(len(report.source_files), 3)
            self.assertEqual(report.groups, [])

            fusion.collect(report, cfg, fusion.COLUMN_SCHEMA)
            self.assertEqual([group.name for group in report.groups], ["Data", "Diğer"])
            data_group, other_group = report.groups
            self.assertEqual([record.values["No."] for record in data_group.records], [2, 3, 1])
            self.assertEqual({record.sheet for record in data_group.records}, {"Data", "data"})
            self.assertEqual(len(data_group.files), 3)
            self.assertEqual(len(data_group.read_files), 3)
            self.assertEqual(data_group.table_count, 3)
            self.assertEqual([record.values["No."] for record in other_group.records], [20])
            self.assertFalse(any(event["Seviye"] == "ERROR" for event in report.events))

            cfg["WORKFLOW_SUMMARY"]["ENABLED"] = False
            fusion.save_report(report, cfg, fusion.COLUMN_SCHEMA)
            import openpyxl
            result = openpyxl.load_workbook(report.output, data_only=True)
            try:
                self.assertIn("Data", result.sheetnames)
                self.assertIn("Diğer", result.sheetnames)
                self.assertNotIn("Grup A", result.sheetnames)
                summary = result[cfg["SUMMARY_SHEET"]]
                self.assertIn("Kaynak Excel sayfası", summary["A3"].value)
                self.assertEqual(summary["A41"].value, "KAYNAK SAYFA BAZINDA ÖZET")
                self.assertEqual(result["Data"]["A8"].value, 2)
            finally:
                result.close()

    def test_import_summary_and_disabled_layout(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source_dir = root / "kaynaklar" / "Grup A"
            source_dir.mkdir(parents=True)
            source = source_dir / "source.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["No.", "Entity", "Panel", "Open/Closed\nAçıklama", "My Company"])
            for row in [(1, "E", "P", " OpEn: uzun açıklama", None),
                        (2, "E", "P", "OPEN\nclosed daha sonra", "yanıt"),
                        (3, "E", "P", "Closed: açıklama", None),
                        (4, "E", "P", "CLOSED", "yanıt"),
                        (5, "E", "P", "pending open", None),
                        (6, "E", "P", "open", "=A2"),
                        ("total", None, None, "open", None)]:
                sheet.append(row)
            missing = workbook.create_sheet("Eksik şirket")
            missing.append(["No.", "Entity", "Panel", "Open / Closed"])
            missing.append([7, "E", "P", "open"])
            workbook.save(source)
            workbook.close()
            original_bytes = source.read_bytes()

            cfg = copy.deepcopy(fusion.CONFIG)
            fusion.validate(cfg, fusion.COLUMN_SCHEMA)
            report = fusion.Report(root / "kaynaklar", root / "report.xlsx")
            fusion.discover(report, cfg)
            fusion.collect(report, cfg, fusion.COLUMN_SCHEMA)
            report.groups.append(fusion.Group("Boş grup"))
            self.assertEqual(len(report.groups[0].records), 7)
            self.assertEqual(len(report.quarantine), 1)
            self.assertFalse(any(event["Seviye"] == "ERROR" for event in report.events))
            self.assertEqual(report.groups[0].records[0].values["Open / Closed"], "OpEn: uzun açıklama")

            for enabled in (True, False):
                with self.subTest(enabled=enabled):
                    cfg["WORKFLOW_SUMMARY"]["ENABLED"] = enabled
                    cfg["OVERWRITE"] = True
                    fusion.save_report(report, cfg, fusion.COLUMN_SCHEMA)
                    result = openpyxl.load_workbook(report.output, data_only=True)
                    try:
                        summary = result[cfg["SUMMARY_SHEET"]]
                        if enabled:
                            self.assertEqual(summary["A12"].value, "AÇIK / KAPALI VE İŞ SIRASI")
                            columns = (1, 6, 8, 10, 13, 16, 19)
                            self.assertEqual([summary.cell(14, col).value for col in columns],
                                             ["GENEL TOPLAM", 4, 2, 1, 1, 1, 2])
                            self.assertEqual([summary.cell(15, col).value for col in columns],
                                             ["Grup A", 4, 2, 1, 1, 1, 2])
                            self.assertEqual([summary.cell(16, col).value for col in columns],
                                             ["Boş grup", 0, 0, 0, 0, 0, 0])
                        else:
                            self.assertIsNone(summary["A12"].value)
                        offset = len(report.groups) + 6 if enabled else 0
                        self.assertEqual(summary.cell(41 + offset, 1).value, "KLASÖR BAZINDA ÖZET")
                        self.assertEqual(summary["A6"].value, 7)
                        self.assertEqual(result["_Rapor Verisi"].sheet_state, "hidden")
                        self.assertEqual(len(summary._charts), 4)
                        self.assertEqual(summary._charts[0].anchor._from.row, 12 + offset)
                        self.assertEqual(len(result["Grup A"].tables), 1)
                    finally:
                        result.close()
            self.assertEqual(source.read_bytes(), original_bytes)


if __name__ == "__main__":
    unittest.main()
